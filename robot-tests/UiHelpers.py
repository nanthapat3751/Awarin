"""ตัวช่วยของชุดหน้าจอที่เขียนใน .robot แล้วอ่านยาก: ตั้งค่าเบราว์เซอร์ (ขนาดจอ/โฟลเดอร์ดาวน์โหลด/Edge)
รอไฟล์ดาวน์โหลด แก้ไฟล์ Excel ที่ส่งออก และรหัส 2FA"""
from __future__ import annotations

import base64
import hashlib
import hmac
import shutil
import struct
import time
from pathlib import Path


class UiHelpers:
    ROBOT_LIBRARY_SCOPE = "GLOBAL"

    def browser_options(self, browser: str = "chrome", width: int = 1366, height: int = 900,
                        download_dir: str | None = None, mobile: bool = False):
        """ตัวเลือกเบราว์เซอร์แบบไม่มีหน้าต่าง — ส่งให้ Open Browser ผ่าน options=
        download_dir = โหลดไฟล์ลงโฟลเดอร์นี้เงียบ ๆ (ไม่ถาม) · mobile = จำลองมือถือ (UA + จอสัมผัส)"""
        from selenium import webdriver

        opts = webdriver.EdgeOptions() if browser.lower() == "edge" else webdriver.ChromeOptions()
        opts.add_argument("--headless=new")
        opts.add_argument(f"--window-size={int(width)},{int(height)}")
        opts.add_argument("--lang=th-TH")
        if mobile:
            opts.add_experimental_option("mobileEmulation", {
                "deviceMetrics": {"width": int(width), "height": int(height), "pixelRatio": 3.0, "touch": True},
                "userAgent": "Mozilla/5.0 (Linux; Android 14; Pixel 8) AppleWebKit/537.36 (KHTML, like Gecko) "
                             "Chrome/155.0 Mobile Safari/537.36",
            })
        if download_dir:
            Path(download_dir).mkdir(parents=True, exist_ok=True)
            opts.add_experimental_option("prefs", {
                "download.default_directory": str(Path(download_dir).resolve()),
                "download.prompt_for_download": False,
                "safebrowsing.enabled": True,
            })
        return opts

    def empty_directory(self, path: str) -> None:
        shutil.rmtree(path, ignore_errors=True)
        Path(path).mkdir(parents=True, exist_ok=True)

    def wait_for_download(self, directory: str, suffix: str, timeout: float = 30) -> str:
        """รอจนมีไฟล์นามสกุลนี้ดาวน์โหลดเสร็จในโฟลเดอร์ (ไม่นับไฟล์ .crdownload ที่ยังโหลดไม่จบ)"""
        until = time.time() + float(timeout)
        while time.time() < until:
            done = [p for p in Path(directory).glob(f"*{suffix}") if p.stat().st_size > 0]
            partial = list(Path(directory).glob("*.crdownload"))
            if done and not partial:
                return str(max(done, key=lambda p: p.stat().st_mtime))
            time.sleep(0.5)
        raise AssertionError(f"ไม่มีไฟล์ {suffix} ดาวน์โหลดเสร็จใน {directory} ภายใน {timeout} วินาที")

    def file_starts_with_bom(self, path: str) -> bool:
        return Path(path).read_bytes()[:3] == b"\xef\xbb\xbf"

    def read_text_file(self, path: str) -> str:
        return Path(path).read_text(encoding="utf-8-sig")

    def set_boq_quantity(self, src: str, dst: str, item_name: str, quantity: float) -> int:
        """แก้ "ปริมาณ" ของบรรทัด BOQ ชื่อ item_name ในไฟล์ที่ส่งออกจากระบบ แล้วบันทึกเป็นไฟล์ใหม่
        (คงคอลัมน์ซ่อนรหัสบรรทัดและชีตซ่อน _systemboq ไว้ ระบบจึงรู้ว่าเป็นการแก้บรรทัดเดิม) · คืนจำนวนแถวที่แก้"""
        import openpyxl

        wb = openpyxl.load_workbook(src)
        changed = 0
        for ws in wb.worksheets:
            if ws.sheet_state != "visible":
                continue
            header_row = None
            for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20)):
                values = [c.value for c in row]
                if "รายการ" in values and "ปริมาณ" in values:
                    header_row = row[0].row
                    name_col = values.index("รายการ") + 1
                    qty_col = values.index("ปริมาณ") + 1
                    break
            if header_row is None:
                continue
            for r in range(header_row + 1, ws.max_row + 1):
                if ws.cell(r, name_col).value == item_name:
                    ws.cell(r, qty_col).value = quantity
                    changed += 1
        wb.save(dst)
        return changed

    def xlsx_summary(self, path: str) -> dict:
        """ชื่อชีตที่มองเห็น ชีตที่ซ่อน และข้อความทุกช่องในชีตที่มองเห็น (ไว้ตรวจว่าไฟล์ที่ส่งออกมีของครบ)"""
        import openpyxl

        wb = openpyxl.load_workbook(path)
        visible = [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"]
        hidden = [ws.title for ws in wb.worksheets if ws.sheet_state != "visible"]
        cells = []
        for ws in wb.worksheets:
            if ws.sheet_state == "visible":
                for row in ws.iter_rows(values_only=True):
                    cells.extend(str(v) for v in row if v not in (None, ""))
        return {"visible": visible, "hidden": hidden, "text": "\n".join(cells)}

    def totp(self, secret: str) -> str:
        """รหัส 6 หลักแบบแอป Authenticator (RFC 6238, 30 วินาที, SHA-1)"""
        key = base64.b32decode(secret.upper() + "=" * (-len(secret) % 8))
        h = hmac.new(key, struct.pack(">Q", int(time.time()) // 30), hashlib.sha1).digest()
        o = h[-1] & 15
        return str((struct.unpack(">I", h[o: o + 4])[0] & 0x7FFFFFFF) % 1_000_000).zfill(6)

    def seconds_into_totp_step(self) -> int:
        """วินาทีที่ผ่านไปในช่วง 30 วินาทีปัจจุบัน — ถ้าใกล้หมดช่วง ให้รอช่วงใหม่ก่อนพิมพ์รหัส (กันรหัสเปลี่ยนกลางทาง)"""
        return int(time.time()) % 30
