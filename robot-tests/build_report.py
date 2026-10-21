"""
สร้างรายงาน Test Case (Excel) จากผลรัน Robot Framework

    python -X utf8 robot-tests/build_report.py --out "D:/ปี4_1/ปัญหาพิเศษ/SystemBOQ-Test-Report.xlsx"

อ่าน robot-tests/results/output.xml แล้วเขียน:
  สรุปผล · แผนการทดสอบ · กลยุทธ์การทดสอบ · Test Cases · Defects
และคัดลอกหลักฐาน (ภาพหน้าจอ + report.html/log.html) ไปไว้ข้างไฟล์ Excel ในโฟลเดอร์ testcase-evidence/

ข้อมูลของแต่ละเคส (Scenario/Pre/Steps/Data/Expected) อ่านจาก [Documentation] ในไฟล์ .robot
ที่เดียว — แก้เคสที่ไฟล์ .robot แล้วสร้างรายงานใหม่ ไม่ต้องแก้ที่นี่
"""
from __future__ import annotations

import argparse
import re
from collections import Counter
import shutil
from datetime import datetime
from pathlib import Path

import io

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.utils.units import pixels_to_EMU
from PIL import Image as PILImage
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink
from robot.api import ExecutionResult

HERE = Path(__file__).resolve().parent
RESULTS = HERE / "results"

ENVIRONMENT = "DEV (localhost) · Chrome 155 headless · Windows 11"
# โมดูลที่ไม่ได้รันบน Chrome — ช่อง Environment ของเคสในโมดูลนี้ต้องบอกเบราว์เซอร์ที่ใช้จริง
MODULE_ENVIRONMENT = {"Cross Browser": "DEV (localhost) · Microsoft Edge 155 headless · Windows 11"}
ENVIRONMENT_FULL = "DEV (localhost · ฐานข้อมูลทดสอบ) · Windows 11 · Chrome 155 (headless) · 1366×900"
TESTER = "Robot Framework 7.5 (อัตโนมัติ)"

# ข้อบกพร่องที่พบ — ขึ้นในรายงานเมื่อ (ก) เคสนั้นไม่ผ่านในรอบนี้ = Open หรือ (ข) ไม่ผ่านในรอบแรก
# (testcase-evidence-round1/output.xml) แต่รอบนี้ผ่านแล้ว = Closed พร้อมผลก่อน/หลังแก้ · ตัดสินจากผลรันจริงทั้งคู่
DEFECTS = {
    "TC_CUS_ACCESS_002": {
        "id": "BUG-001",
        "severity": "Medium",
        "priority": "P2",
        "title": "ลูกค้าเข้าสู่ระบบแล้วไม่กลับไปหน้าที่ตั้งใจเปิด (ถูกพาไปหน้าแรกเสมอ)",
        "impact": "ลิงก์ในอีเมลแจ้งเตือน (เช่น ใบเสนอราคาใหม่) พาลูกค้าไปหน้าเข้าสู่ระบบ "
        "แต่หลังเข้าสู่ระบบลูกค้าถูกพาไปหน้าแรก ต้องหาใบเสนอราคาเองอีกรอบ",
        "fix": "ProtectedRoute ส่งหน้าที่ตั้งใจเปิดไปกับ <Navigate state={{ from: location }}> "
        "แล้ว LoginForm พากลับไปที่ from หลังเข้าสู่ระบบสำเร็จ (รับเฉพาะ path ภายในเว็บ กัน open redirect)",
        "where": "client-app/src/components/ProtectedRoute.tsx · client-app/src/components/LoginForm.tsx",
        "change": "เพิ่ม lib/return-to.ts (returnPathFrom: รับเฉพาะ path ภายใน ไม่รับ //… /\\… URL ภายนอก /login /register) · "
        "ProtectedRoute ส่ง state.from · LoginForm และ AuthDuoPage พากลับไปที่ from · ตรวจตัวกรองด้วย 12 กรณี (รวม open redirect) ผ่านครบ",
    },
    "TC_PROJ_002": {
        "id": "BUG-002",
        "severity": "Low",
        "priority": "P3",
        "title": "กด \"สร้าง\" โครงการโดยไม่กรอกชื่อ ระบบเงียบ ไม่แจ้งอะไร",
        "impact": "ไม่สร้างโครงการที่ไม่มีชื่อ (ถูกต้อง) แต่ผู้ใช้ไม่รู้ว่าทำไมกดแล้วไม่เกิดอะไรขึ้น "
        "ต่างจากฟอร์ม BOQ ที่บอกว่าขาดช่องไหน",
        "fix": "ใน submitCreate ถ้าชื่อว่างให้ toast.error(\"กรอกชื่อโครงการ\") แทน return เงียบ ๆ",
        "where": "staff-app/src/pages/ProjectsPage.tsx (submitCreate)",
        "change": "submitCreate แจ้ง toast.error(\"กรอกชื่อโครงการก่อน\") เมื่อชื่อว่าง แทน return เงียบ ๆ (ยังไม่ส่งคำขอสร้างโครงการเหมือนเดิม)",
    },
}

REMARKS = {
    "TC_STAFF_LOGIN_007": "ครั้งที่ 5 ยังตอบ \"รหัสผิด\" ตามออกแบบ (ไม่บอกผู้โจมตีว่าล็อกแล้ว) · ล็อกเริ่มมีผลครั้งที่ 6 "
    "· ระบบแจ้ง OWNER/ADMIN ทางกระดิ่ง · seed ปลดล็อกบัญชีให้ทุกรอบ",
    "TC_STAFF_LOGIN_003": "กันการไล่เดาชื่อผู้ใช้ (account enumeration)",
    "TC_BOQ_002": "ตอบคำถามเดิม: ปุ่ม \"บันทึก\" ปิดฟอร์ม ส่วน \"บันทึก + เพิ่มต่อ\" ตั้งใจให้ฟอร์มเปิดค้าง (ดู TC_BOQ_004)",
    "TC_BOQ_003": "ใช้ค่าดำเนินการ+กำไรมาตรฐาน 10% และ VAT 7% ของฐานข้อมูลทดสอบ",
    "TC_CUS_LOGIN_003": "ข้อสังเกต: อีเมลของลูกค้าอยู่ใน URL (?email=…) ซึ่งขัดกับหลักใน RegisterForm ที่ห้ามข้อมูลส่วนตัวอยู่ใน URL "
    "— ไม่นับเป็นบัค ควรพิจารณา",
    "TC_CUS_ACCESS_002": "ฝั่งหลังบ้านแก้แบบเดียวกันแล้ว 19 ต.ค. — ทดสอบที่ TC_STAFF_ACCESS_004",
    "TC_STAFF_ACCESS_004": "เพิ่มรอบที่ 4 หลังแก้ staff-app ตามข้อสังเกตจาก BUG-001 (พบจากการอ่านโค้ด ไม่ได้ออก Defect ID เพราะไม่มีผลรันก่อนแก้) "
    "· แก้ทางเข้าสู่ระบบแบบ 2FA ด้วย แต่ชุดนี้ไม่มีบัญชีเปิด 2FA จึงยังไม่ได้ทดสอบทางนั้น",
    "TC_REG_001": "ทุกเคสสมัครสมาชิกหยุดที่การตรวจในหน้าจอ ไม่สร้างบัญชีและไม่ส่งอีเมล",
    "TC_PAY_CUS_003": "สลิปจำลอง (SLIP_PROVIDER=mock) — ระบบอ่านยอด/อ้างอิง/วันที่/บัญชีจากชื่อไฟล์ ไม่เรียก SlipOK จริง "
    "จึงยืนยันตรรกะตรวจสลิปของระบบ ไม่ได้ยืนยันการต่อกับ SlipOK",
    "TC_PAY_STAFF_001": "ต่อจากการแจ้งโอนใน TC_PAY_CUS_003 (หน้าบ้านรันก่อนหลังบ้านเสมอ) — ถ้าไม่มีการแจ้ง เคสจะถูกข้าม (Block)",
    "TC_PAY_STAFF_003": "คำเตือนเป็นแบบเตือน ไม่ห้าม (ระบบรับเงินเกินได้โดยตั้งใจ) — เคสนี้กดกลับไปแก้ ไม่ได้บันทึก",
    "TC_PAY_STAFF_005": "บัญชี OWNER ยืนยันรายการที่ตัวเองบันทึกได้ (บริษัทคนเดียว) — ADMIN ยืนยันของตัวเองไม่ได้ ยังไม่ได้ทดสอบกรณีนั้น",
}

# ชื่อของรอบที่เก็บไว้ (testcase-evidence-round<N>) — รอบปัจจุบันตั้งด้วย --label
ROUND_LABELS = {1: "ก่อนแก้บัค", 2: "หลังแก้ BUG-001/002", 3: "หลังเพิ่ม/ปรับส่วนชำระเงิน", 4: "หลังแก้หลังบ้านกลับหน้าเดิม",
                5: "เพิ่มเคสส่วนชำระเงิน", 6: "ปิดช่องว่าง UI-GAP 16 ข้อ"}

# โมดูลที่เพิ่มในรอบปิดช่องว่าง UI-GAP-01–16 (ชื่อตรงกับ module: ในไฟล์ .robot) — ใช้เขียนขอบเขตในแผนการทดสอบ
GAP_MODULES = [
    ("Quote", "ใบเสนอราคา"), ("Contract", "สัญญา"), ("Billing", "วางบิล/ยกเลิกใบ"),
    ("Customer Payment History", "ประวัติการชำระของลูกค้า"), ("Contact Request", "ฟอร์มติดต่อ"),
    ("Lead To Project", "คำขอ → โครงการ"), ("Role Visibility", "สิทธิ์ ADMIN บนหน้าจอ"), ("Two Factor Login", "2FA"),
    ("Customer Forgot Password", "ลืมรหัสลูกค้า"), ("Staff Forgot Password", "ลืมรหัสพนักงาน"),
    ("Register Success", "สมัครจนเข้าพอร์ทัล"), ("Customer File", "แฟ้มลูกค้า"), ("BOQ Excel", "Excel BOQ"),
    ("Mobile Layout", "มือถือ 390 px"), ("iPad Layout", "iPad"), ("Trash", "ถังขยะ"), ("Notifications", "กระดิ่ง"),
    ("Cross Browser", "Edge"), ("Keyboard Access", "คีย์บอร์ด/ป้ายอ่านหน้าจอ"), ("Stored XSS UI", "ชื่อที่มี HTML"),
]

STATUS_MAP = {"PASS": "Pass", "FAIL": "Fail", "SKIP": "Block", "NOT RUN": "Untested"}

FONT = "Tahoma"  # มีตัวอักษรไทยในตัว (Arial ไม่มี — Excel จะสลับฟอนต์เองทีละตัวอักษร)
NAVY = "000099"
thin = Side(style="thin", color="BFC5D2")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def est_lines(text, width: float) -> int:
    """จำนวนบรรทัดโดยประมาณเมื่อห่อข้อความในคอลัมน์กว้าง width (สระ/วรรณยุกต์ไทยไม่กินที่ จึงคูณ 1.25)"""
    per_line = max(width * 1.25, 1)
    return sum(max(1, -(-len(line) // int(per_line))) for line in str(text).split("\n"))


BASE_DIR = Path(".")  # โฟลเดอร์ของไฟล์ Excel — ตั้งใน main() ใช้หาไฟล์ภาพจาก path สัมพัทธ์
THUMB_W = 320  # ความกว้างภาพหน้าจอที่แปะในเซลล์ (px)
CELL_PAD = 4


def embed_screenshot(ws, rel_path: str, row: int, col: int, width_px: int = THUMB_W) -> int:
    """แปะภาพหน้าจอลงในเซลล์ (ย่อและบีบเป็น JPEG ก่อน ไม่งั้นไฟล์บวมหลายสิบ MB) คืนความสูงภาพเป็น px

    ยึดทั้งสองมุมไว้ในเซลล์เดียวกัน (twoCell) — กรองแถวด้วยตัวกรองของ Excel แล้วภาพหายไปพร้อมแถว ไม่ลอยค้างทับแถวอื่น
    """
    src = BASE_DIR / rel_path
    if not rel_path or not src.exists():
        return 0
    with PILImage.open(src) as im:
        im = im.convert("RGB")
        height_px = round(im.height * width_px / im.width)
        im = im.resize((width_px * 2, height_px * 2), PILImage.LANCZOS)  # ความละเอียด 2 เท่า ซูมแล้วยังอ่านออก
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=80, optimize=True)
    buf.seek(0)
    img = XLImage(buf)
    img.width, img.height = width_px, height_px
    pad = pixels_to_EMU(CELL_PAD)
    img.anchor = TwoCellAnchor(
        editAs="twoCell",
        _from=AnchorMarker(col=col - 1, colOff=pad, row=row - 1, rowOff=pad),
        to=AnchorMarker(col=col - 1, colOff=pad + pixels_to_EMU(width_px), row=row - 1, rowOff=pad + pixels_to_EMU(height_px)),
    )
    ws.add_image(img)
    return height_px


def image_col_width(width_px: int = THUMB_W) -> float:
    """ความกว้างคอลัมน์ (หน่วยตัวอักษรของ Excel ≈ 7px) ที่พอดีกับภาพ + ขอบ"""
    return round((width_px + CELL_PAD * 2 + 5) / 7, 1)


def f(size=10, bold=False, color="1F2937", italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


# ---------------------------------------------------------------- อ่านผล robot
def parse_doc(doc: str) -> dict[str, str]:
    out: dict[str, str] = {}
    for line in doc.splitlines():
        m = re.match(r"\s*(Scenario|Pre|Steps|Data|Expected):\s*(.*)", line)
        if m:
            out[m.group(1)] = m.group(2).strip()
    return out


def as_lines(text: str, sep: str) -> str:
    return "\n".join(p.strip() for p in text.split(sep) if p.strip())


def read_round(path: Path, include=None) -> tuple[dict, dict] | tuple[None, None]:
    """ผลรันรอบก่อน (เก็บไว้ที่ testcase-evidence-round1/) — ใช้แสดงประวัติและผลก่อนแก้บัค"""
    if not path.exists():
        return None, None
    res = ExecutionResult(str(path))
    tests: dict[str, object] = {}

    def walk(s):
        for t in s.tests:
            tests[t.name.split(" ", 1)[0]] = t
        for c in s.suites:
            walk(c)

    # รอบหลัง ๆ มีชุด API รวมอยู่ในไฟล์เดียวกัน — นับเฉพาะชุดที่รายงานนี้สนใจ (รอบเก่าไม่มีชุด api อยู่แล้ว)
    picked = [c for c in res.suite.suites if include is None or c.name.lower() in include] or [res.suite]
    for c in picked:
        walk(c)
    status = [t.status for t in tests.values()]
    start = min((c.start_time for c in picked if c.start_time), default=res.suite.start_time)
    stats = {"start": start, "total": len(status), "pass": status.count("PASS"),
             "fail": status.count("FAIL"), "skip": status.count("SKIP")}
    return tests, stats


# ชุดบนสุดใน output.xml (รวมด้วย rebot) → ชื่อส่วนของระบบในรายงาน
LAYERS = {"frontend": "หน้าบ้าน (client-app)", "backend": "หลังบ้าน (staff-app)", "api": "API (server)"}
UI_SUITES = ("frontend", "backend")
EVIDENCE_PLACEHOLDER = "ไม่มีภาพ"
# หมายเหตุท้ายเคสที่รอบแรกไม่ผ่านแล้วแก้ผ่าน (รายงาน API ไม่มีภาพ จึงเปลี่ยนข้อความได้)
FIXED_NOTE = "ภาพก่อน/หลังแก้อยู่ในชีต Defects"
# ข้อบกพร่องที่พบในรอบหลัง (เคสเพิ่งมีในรอบนั้น รอบที่ 1 ไม่มี) — {รหัสเคส: (โฟลเดอร์ผลรอบที่พบ, "รอบที่ N")}
# ผลก่อนแก้/วันที่พบอ่านจากโฟลเดอร์นั้นแทน testcase-evidence-round1 · ว่าง = ทุกบัคพบในรอบที่ 1 (รายงาน UI)
FOUND_IN: dict[str, tuple[Path, str]] = {}


def collect(round1_dir: Path, include=UI_SUITES, evidence_dir: str = "testcase-evidence") -> tuple[list[dict], dict]:
    """include = ชุดบนสุดที่จะเอามาลงรายงาน (รายงาน UI = หน้าบ้าน+หลังบ้าน · รายงาน API = api)"""
    result = ExecutionResult(str(RESULTS / "output.xml"))
    round1, round1_stats = read_round(round1_dir / "output.xml", include)
    rows: list[dict] = []
    later_rounds = {d: read_round(d / "output.xml", include)[0] or {} for d, _ in FOUND_IN.values()}

    def walk(suite, layer):
        for t in suite.tests:
            tc_id = t.name.split(" ", 1)[0]
            title = t.name.split(" ", 1)[1] if " " in t.name else t.name
            doc = parse_doc(t.doc)
            module = next((tag.split(":", 1)[1] for tag in t.tags if tag.startswith("module:")), suite.name)
            status = STATUS_MAP.get(t.status, "Untested")
            known = DEFECTS.get(tc_id)
            found_dir, found_label = FOUND_IN.get(tc_id, (round1_dir, "รอบที่ 1"))
            before = (later_rounds[found_dir] if tc_id in FOUND_IN else (round1 or {})).get(tc_id)
            fixed = bool(known and status == "Pass" and before is not None and before.status == "FAIL")
            defect = known if (status == "Fail" or fixed) else None
            if status == "Pass":
                actual = t.message or "ตรงตามผลที่คาดไว้"
            elif status == "Fail":
                # ข้อความล้มบางเคสพ่วงข้อมูลทดสอบยาวมาก (เช่น คำค้น 5,000 ตัว) — ตัดให้อ่านได้ ข้อความเต็มอยู่ใน log.html
                actual = t.message if len(t.message) <= 400 else t.message[:400] + " … (ข้อความเต็มใน log.html)"
            else:
                actual = t.message or "ไม่ได้รัน"
            remarks = REMARKS.get(tc_id, "")
            if status == "Fail" and not defect:
                remarks = (remarks + " · " if remarks else "") + "ไม่ผ่านโดยยังไม่ได้วิเคราะห์ — ตรวจว่าเป็นบัคหรือสคริปต์ก่อนออก Defect ID"
            if fixed:
                remarks = (f"{found_label} ไม่ผ่าน ({known['id']}) → แก้แล้ว ทดสอบซ้ำผ่าน · "
                           f"{FIXED_NOTE}") + (" · " + remarks if remarks else "")
            evidence = RESULTS / "evidence" / f"{tc_id}.png"
            before_evidence = found_dir / "evidence" / f"{tc_id}.png"
            rows.append(
                {
                    "id": tc_id,
                    "layer": layer,
                    "module": module,
                    "title": doc.get("Scenario") or title,
                    "pre": doc.get("Pre", ""),
                    "steps": as_lines(doc.get("Steps", ""), "|"),
                    "data": as_lines(doc.get("Data", ""), "·"),
                    "expected": doc.get("Expected", ""),
                    "actual": actual,
                    "status": status,
                    "defect": defect["id"] if defect else "",
                    "severity": f'{defect["severity"]} / {defect["priority"]}' if defect else "",
                    "date": t.start_time,
                    "evidence": f"{evidence_dir}/evidence/{tc_id}.png" if evidence.exists() else "",
                    "remarks": remarks,
                    "defect_info": defect,
                    "defect_state": "Closed" if fixed else ("Open" if defect else ""),
                    "before_actual": before.message if fixed else "",
                    "before_date": before.start_time if fixed else None,
                    "before_evidence": (f"{found_dir.name}/evidence/{tc_id}.png"
                                        if fixed and before_evidence.exists() else ""),
                }
            )
        for child in suite.suites:
            walk(child, layer)

    picked = [c for c in result.suite.suites if c.name.lower() in include]
    for top in picked:
        walk(top, LAYERS.get(top.name.lower(), top.name))
    # ผลที่รวมด้วย rebot (แต่ละชุดคนละรอบ) — ชุดบนสุดไม่มีเวลาของตัวเอง ใช้ของชุดที่เลือก
    starts = [s.start_time for s in picked if s.start_time]
    ends = [s.end_time for s in picked if s.end_time]
    start, end = min(starts), max(ends)
    meta = {
        "start": start,
        "end": end,
        "elapsed": end - start,
        "round1": round1_stats,
    }
    return rows, meta


# ---------------------------------------------------------------- ชีตข้อความ (แผน/กลยุทธ์)
def write_doc_sheet(ws, title: str, subtitle: str, sections: list[tuple[str, list]]):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 48
    ws.column_dimensions["D"].width = 48
    ws.column_dimensions["E"].width = 34
    ws["B2"] = title
    ws["B2"].font = f(16, True, NAVY)
    ws["B3"] = subtitle
    ws["B3"].font = f(10, color="6B7280")
    r = 5
    for heading, body in sections:
        ws.cell(r, 2, heading).font = f(12, True, NAVY)
        ws.cell(r, 2).border = Border(bottom=Side(style="medium", color=NAVY))
        for c in range(3, 6):
            ws.cell(r, c).border = Border(bottom=Side(style="medium", color=NAVY))
        r += 1
        if body and isinstance(body[0], tuple) and body[0][0] == "__table__":
            header = body[0][1]
            for i, h in enumerate(header):
                cell = ws.cell(r, 2 + i, h)
                cell.font = f(10, True, "FFFFFF")
                cell.fill = fill(NAVY)
                cell.alignment = CENTER
                cell.border = BORDER
            r += 1
            for row in body[1:]:
                height = 1
                for i, v in enumerate(row):
                    cell = ws.cell(r, 2 + i, v)
                    cell.font = f(10, bold=(i == 0))
                    cell.alignment = WRAP_TOP
                    cell.border = BORDER
                    width = ws.column_dimensions[get_column_letter(2 + i)].width * 0.9
                    height = max(height, est_lines(v, width))
                ws.row_dimensions[r].height = 15 * height + 5
                r += 1
        else:
            for item in body:
                if isinstance(item, tuple):
                    k, v = item
                    ws.cell(r, 2, k).font = f(10, True)
                    ws.cell(r, 2).alignment = WRAP_TOP
                    ws.cell(r, 3, v).font = f(10)
                    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
                    ws.cell(r, 3).alignment = WRAP_TOP
                    lines = max(est_lines(v, 125), est_lines(k, 27))
                    ws.row_dimensions[r].height = 15 * lines + 5
                else:
                    ws.cell(r, 2, "•  " + item).font = f(10)
                    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
                    ws.cell(r, 2).alignment = WRAP_TOP
                    ws.row_dimensions[r].height = 15 * est_lines("•  " + item, 155) + 5
                r += 1
        r += 1
    return r


def plan_sections(rows, meta):
    total = len(rows)
    fe = sum(1 for x in rows if x["layer"].startswith("หน้าบ้าน"))
    be = total - fe
    run_day = meta["start"].strftime("%d/%m/%Y")
    cnt = Counter(x["module"] for x in rows)
    history = meta.get("history") or []
    # รันซ้ำหลังรอบแก้บัค (รอบที่ 3 เป็นต้นไป + รอบนี้) — ขั้น 3.1, 3.2, …
    reruns = [(h["label"], h["start"].strftime("%d/%m/%Y"), h["pass"], h["total"]) for h in history if h["n"] >= 3]
    if len(history) >= 2:
        reruns.append((meta["label"], run_day, sum(1 for x in rows if x["status"] == "Pass"), len(rows)))
    rerun_rows = [(f"3.{i}", f"รันซ้ำทั้งชุด — {label}", "ยืนยันว่าส่วนเดิมไม่พังหลังแก้โค้ด", f"เสร็จ {day} — ผ่าน {p}/{t}")
                  for i, (label, day, p, t) in enumerate(reruns, start=1)]
    first_day = history[0]["start"].strftime("%d/%m/%Y") if history else run_day
    # รอบที่แก้บัคเสร็จ = รอบที่ 2 ถ้าเก็บไว้ (ไม่งั้นคือรอบนี้)
    fixed_round = next((h for h in history if h["n"] == 2), None)
    closed = sorted(x["defect"] for x in rows if x["defect_state"] == "Closed")
    still_open = sorted(x["defect"] for x in rows if x["defect_state"] == "Open")
    passed = sum(1 for x in rows if x["status"] == "Pass")
    if still_open:
        fix_status = "ค้าง " + ", ".join(still_open)
    elif closed and fixed_round:
        fix_status = (f'เสร็จ {fixed_round["start"]:%d/%m/%Y} — แก้ {", ".join(closed)} แล้ว '
                      f'รันซ้ำทั้งชุดผ่าน {fixed_round["pass"]}/{fixed_round["total"]}')
    elif closed:
        fix_status = f"เสร็จ {run_day} — แก้ {', '.join(closed)} แล้ว รันซ้ำทั้งชุดผ่าน {passed}/{total}"
    else:
        fix_status = "ไม่มีข้อบกพร่องให้แก้"
    return [
        ("1. วัตถุประสงค์", [
            "ยืนยันว่าการทำงานหลักของระบบ SystemBOQ ทั้งหน้าบ้าน (เว็บลูกค้า) และหลังบ้าน (ระบบพนักงาน) ทำงานตามที่ออกแบบ ก่อนสแกน SonarQube และก่อนขึ้นระบบจริง",
            "ทดสอบแบบอัตโนมัติ (Automated Testing) ด้วย Robot Framework เพื่อรันซ้ำได้ทุกครั้งที่แก้โค้ด (regression)",
            "บันทึกผลเป็นหลักฐาน (ภาพหน้าจอ + รายงาน Robot) และติดตามข้อบกพร่องที่พบด้วยรหัส BUG-###",
        ]),
        ("2. ขอบเขตการทดสอบ", [
            ("อยู่ในขอบเขต (รอบนี้)",
             f'หน้าบ้าน {fe} เคส: หน้าสาธารณะ {cnt["Public Website"]} · '
             f'เข้าสู่ระบบลูกค้า/สิทธิ์เข้าถึง {cnt["Customer Login"] + cnt["Customer Access Control"]} · '
             f'ตรวจข้อมูลฟอร์มสมัครสมาชิก {cnt["Customer Registration"]}\n'
             f'หลังบ้าน {be} เคส: เข้าสู่ระบบพนักงาน/ล็อกบัญชี {cnt["Staff Login"]} · สิทธิ์เข้าถึง {cnt["Staff Access Control"]} · '
             f'สร้างโครงการ + BOQ + การคิดราคา {cnt["Project"] + cnt["BOQ"] + cnt["BOQ Pricing"]}'
             + (f'\nส่วนชำระเงิน {cnt["Customer Payments"] + cnt["Staff Payments"]} เคส: ลูกค้าแจ้งโอน/ตรวจสลิป/ยกเลิก {cnt["Customer Payments"]} · '
                f'พนักงานยืนยันการแจ้ง/บันทึกรับเงิน/ตรวจทาน/ยืนยันกับบัญชี/ออกใบเสร็จ {cnt["Staff Payments"]}'
                if cnt["Customer Payments"] + cnt["Staff Payments"] else "")
             + (("\nปิดช่องว่าง UI-GAP-01–16 " + str(sum(cnt[m] for m, _ in GAP_MODULES)) + " เคส: "
                 + " · ".join(f"{label} {cnt[m]}" for m, label in GAP_MODULES if cnt[m]))
                if any(cnt[m] for m, _ in GAP_MODULES) else "")),
            ("นอกขอบเขต (รอบถัดไป)",
             ("ต่อ SlipOK จริง · QR พร้อมเพย์ · เงินประกันผลงาน · งานเพิ่ม-ลด (VO) · Firefox (ผู้ใช้ตัดออกจากขอบเขต 21 ต.ค. — ลูกค้าส่วนใหญ่ใช้ Chrome/Edge/Safari) · "
              "มือถือ/iPad เครื่องจริง (Safari) · ประสิทธิภาพและการรับโหลด"
              if any(cnt[m] for m, _ in GAP_MODULES) else
              "ใบเสนอราคา/สัญญา (สร้าง ส่ง ลงนาม) · วางบิล · การสร้าง PDF (ใบแจ้งหนี้ ใบเสร็จ) · ต่อ SlipOK จริง · QR พร้อมเพย์ · เงินประกันผลงาน · "
              "ADMIN ยืนยันรายการของตัวเอง · 2FA · ลืมรหัสผ่าน (ต้องอ่านเมล) · นำเข้า/ส่งออก Excel · "
              "จอมือถือ/iPad · เบราว์เซอร์อื่นนอกจาก Chrome · ประสิทธิภาพและการรับโหลด")),
            ("ทดสอบในระดับอื่นแล้ว",
             "ตรรกะฝั่งเซิร์ฟเวอร์ (คิดราคา ภาษี สิทธิ์ API อีเมล) ครอบด้วย Jest + supertest ใน server/tests — "
             "รายงานนี้เน้นสิ่งที่ผู้ใช้เห็นบนหน้าจอ"),
        ] + ([(f'หมายเหตุรอบที่ {len(meta["history"]) + 1}', meta["note"])] if meta.get("note") else [])),
        ("3. รายการที่ทดสอบ (Test Items)", [("__table__", ["ส่วนของระบบ", "ที่อยู่ (ระบบทดสอบ)", "เทคโนโลยี", "บัญชีทดสอบ"]),
            ("หน้าบ้าน (client-app)", "http://localhost:5184", "React + Vite",
             "robot.customer@example.test (ยืนยันแล้ว) · robot.unverified (ยังไม่ยืนยัน)\nrobot.sales (เดินงานขาย) · robot.forgot (ลืมรหัส) · "
             "robot.xss (ชื่อมี HTML)\nrobot.reg.<รอบ>@example.com (สมัครใหม่ทุกรอบ)"),
            ("หลังบ้าน (staff-app)", "http://localhost:5183", "React + Vite",
             "robot_owner (OWNER) · robot_lock (ทดสอบล็อกบัญชี) · robot_admin (ADMIN)\n"
             "robot_twofa (ADMIN เปิด 2FA) · robot_bell (ADMIN กระดิ่ง) · robot_forgot (ADMIN ลืมรหัส)"),
            ("API (server)", "http://localhost:4100", "Node.js + Express + Prisma", "—"),
            ("ฐานข้อมูล", "DATABASE_URL_TEST", "PostgreSQL (Docker)", "ฐานเดียวกับ Jest แยกจากฐาน dev"),
        ]),
        ("4. สภาพแวดล้อมการทดสอบ", [
            ("เครื่อง/เบราว์เซอร์", "Windows 11 · Google Chrome 155 แบบ headless · ขนาดจอ 1366×900 (มือถือจำลอง 390×844 · iPad 1024×1366) · "
             "Microsoft Edge 155 (เคสเบราว์เซอร์อื่น) · Python 3.13 · Robot Framework 7.5 · SeleniumLibrary 6.9"),
            ("แยกจาก dev", "เปิดชุดทดสอบที่พอร์ต 4100/5183/5184 คู่กับ dev server (4000/5173/5174) ได้โดยไม่ชนกัน"),
            ("อีเมล", "MAIL_TRANSPORT=console — ระบบไม่ส่งอีเมลจริงระหว่างทดสอบ พิมพ์เนื้อเมลลง log แทน (dev server ปกติส่งจริงผ่าน Gmail) · "
             "เคสสมัคร/ลืมรหัสอ่านรหัส 6 หลักและลิงก์จาก log (robot-tests/MailLog.py)"),
            ("สลิป", "SLIP_PROVIDER=mock — อ่านผลจากชื่อไฟล์สลิป ไม่เรียก SlipOK จริง (server/.env ต่อ SlipOK ไว้และโควตามีจำกัด)"),
            ("อื่น ๆ", "RATE_LIMIT_STORE=memory · ปิดสำรองข้อมูลอัตโนมัติ · ไฟล์อัปโหลดลงโฟลเดอร์ชั่วคราว · ปิดเช็ค MX"),
        ]),
        ("5. เกณฑ์เริ่มและเกณฑ์จบ", [
            ("เกณฑ์เริ่ม (Entry)", "ฐานข้อมูลทดสอบ migrate ครบ · seed บัญชีทดสอบสำเร็จ · API /api/health ตอบ 200 · ทั้งสองแอปเปิดได้"),
            ("เกณฑ์จบ (Exit)", "รันครบ 100% ของเคสในขอบเขต · ไม่มีข้อบกพร่องระดับ Critical/High ค้าง · "
             "เคสที่ไม่ผ่านทุกเคสมี Defect ID และวิเคราะห์แล้วว่าเป็นบัคของระบบ ไม่ใช่ของสคริปต์"),
            ("เกณฑ์หยุดชั่วคราว (Suspension)", "API หรือแอปเปิดไม่ขึ้น · seed ล้มเหลว · ติดเพดานเข้าสู่ระบบ 20 ครั้ง/15 นาที (ให้เปิดระบบทดสอบใหม่)"),
        ]),
        ("6. ลำดับงานและกำหนดการ", [("__table__", ["ขั้น", "งาน", "ผลลัพธ์", "สถานะ"]),
            ("1", "เขียนเคสทดสอบอัตโนมัติ (Robot Framework)", "robot-tests/ frontend + backend", f"เสร็จ {first_day}"),
            ("2", "รันทดสอบ + บันทึกผล + ออก Defect ID", "รายงานนี้ + testcase-evidence/", f"เสร็จ {first_day}"),
            ("3", "แก้ข้อบกพร่องแล้วรันซ้ำ (regression)", "สถานะเคสที่ไม่ผ่านเปลี่ยนเป็น Pass", fix_status),
        ] + rerun_rows + [
            ("4", "รัน jest --coverage ฝั่ง server", "server/coverage/lcov.info", "รอดำเนินการ"),
            ("5", "สแกน SonarQube แล้วแก้ตามผล", "Quality Gate ผ่าน", "รอดำเนินการ"),
            ("6", "ขึ้นระบบจริง (deploy)", "ระบบจริง", "รอข้อมูลโดเมน/เซิร์ฟเวอร์"),
            ("7", "Smoke test บนระบบจริง (อ่านอย่างเดียว ไม่ส่งเมล)", "ยืนยันว่าระบบจริงเปิดได้", "รอดำเนินการ"),
        ]),
        ("7. ความเสี่ยงและวิธีรับมือ", [("__table__", ["ความเสี่ยง", "ผลกระทบ", "วิธีรับมือ"]),
            ("ระบบส่งอีเมลจริงระหว่างเทสต์", "ลูกค้า/คนอื่นได้อีเมลทดสอบ",
             "ชุดทดสอบตั้ง MAIL_TRANSPORT=console (พิมพ์ลง log ไม่ส่งออก) และใช้โดเมน example.test/example.com"),
            ("สลิปทดสอบไปกินโควตา SlipOK จริง", "โควตาหมดก่อนใช้งานจริง", "ชุดทดสอบบังคับ SLIP_PROVIDER=mock และล้าง SLIPOK_API_KEY ทั้งใน API และ seed"),
            ("เพดานเข้าสู่ระบบ 20 ครั้ง/15 นาที/ไอพี (ใช้ร่วมกันทั้งพนักงานและลูกค้า)", "รันซ้ำติดกันแล้วเทสต์ล้มแบบสุ่ม · ทั้งชุดเข้าสู่ระบบ 21 ครั้ง เกินเพดานถ้ารันรวดเดียว",
             "test-env.mjs --run แบ่งเป็นกลุ่ม (หน้าบ้านเดิม · หน้าบ้าน UI-GAP · หลังบ้านเดิม · หลังบ้าน UI-GAP 2 กลุ่ม · API) เปิด API ใหม่ก่อนทุกกลุ่ม "
             "(ตัวนับอยู่ในหน่วยความจำ เริ่มที่ศูนย์) แล้วรวมผลเป็นรายงานเดียว"),
            ("หน้าจอโหลดไม่ทันตอนเทสต์ตรวจ", "เคสล้มแบบสุ่มทั้งที่ระบบไม่พัง", "ทุกเคสรอสิ่งที่ตัวเองใช้ก่อนตรวจ (พบที่ TC_PAY_CUS_002 รอบทดลอง แก้แล้ว รันซ้ำ 2 รอบผ่านครบ)"),
            ("สถานะค้างข้ามเคส", "เคสก่อนทิ้ง \"หน้าที่ตั้งใจเปิด\" ไว้ เคสถัดไปเข้าสู่ระบบแล้วไปผิดหน้า",
             "keyword เปิดหน้าเข้าสู่ระบบผ่าน about:blank ก่อนเสมอ (พบในรอบที่ 4 แก้แล้ว)"),
            ("ข้อมูลค้างจากรอบก่อน/จาก Jest", "ผลไม่นิ่ง", "seed รีเซ็ตบัญชีทุกรอบ · ชื่อโครงการต่อท้ายด้วยวันเวลา"),
            ("ข้อความหน้าจอเปลี่ยน", "selector หาไม่เจอ เคสล้มทั้งที่ระบบไม่พัง", "ใช้ id/aria-label ก่อนข้อความ · วิเคราะห์ทุกเคสที่ล้มก่อนออก Defect ID"),
            ("ทดสอบเฉพาะ Chrome บนจอคอม", "บั๊กบนมือถือ/เบราว์เซอร์อื่นหลุด",
             "รอบที่ 6 เพิ่มมือถือจำลอง 390 px · iPad 1024 px · Edge — Firefox ตัดออกจากขอบเขต (ผู้ใช้ตัดสินใจ) · เครื่องจริง/Safari ยังไม่ได้ทดสอบ"),
            ("เคสที่ต้องอ่านเมล (สมัคร/ลืมรหัส)", "ถ้าส่งเมลจริงจะไปถึงคนอื่น / อ่านรหัสจากฐานข้อมูลไม่ได้ (เก็บแต่แฮช)",
             "ระบบทดสอบพิมพ์เมลลง log · MailLog.py อ่านฉบับล่าสุดของอีเมลนั้นจาก log ของ API ที่เปิดอยู่"),
            ("เคสที่ต่อกันเป็นลำดับ (งานขาย → กระดิ่ง · ฟอร์มติดต่อ → คำขอ)", "เคสแรกล้มแล้วเคสหลังล้มตาม",
             "เคสหลังข้าม (Block) พร้อมบอกว่ารอเคสไหน · วิเคราะห์จากเคสแรกที่ล้มเสมอ"),
        ]),
        ("8. บทบาทและสิ่งที่ส่งมอบ", [
            ("ผู้รับผิดชอบการทดสอบ", "(กรอกชื่อ — ช่องสีเหลืองในชีตสรุปผล)"),
            ("ผู้รันทดสอบ", "Robot Framework (อัตโนมัติ) — สั่งรันด้วย node robot-tests/test-env.mjs --run"),
            ("สิ่งที่ส่งมอบ", "ไฟล์ Excel นี้ (แผน · กลยุทธ์ · Test Cases · Defects · สรุปผล) + testcase-evidence/ (ภาพหน้าจอทุกเคส, report.html, log.html, output.xml) + สคริปต์ robot-tests/ ในโปรเจกต์"),
        ]),
    ]


def strategy_sections():
    return [
        ("1. ระดับการทดสอบ (Test Levels)", [("__table__", ["ระดับ", "เครื่องมือ", "ทดสอบอะไร", "รันเมื่อไร"]),
            ("Static Analysis", "SonarQube", "กลิ่นโค้ด ช่องโหว่ ความซับซ้อน หนี้ทางเทคนิค (ไม่ต้องรันระบบ)", "หลังรันเทสต์ + coverage ก่อน deploy"),
            ("Unit / Integration", "Jest + supertest", "ตรรกะเซิร์ฟเวอร์ คิดราคา ภาษี สิทธิ์ API การส่งเมล (outbox)", "ทุกครั้งที่แก้โค้ด server · ส่ง coverage ให้ SonarQube"),
            ("End-to-End (UI)", "Robot Framework + SeleniumLibrary", "สิ่งที่ผู้ใช้เห็นและกดจริงบนเบราว์เซอร์ ทั้งหน้าบ้านและหลังบ้าน", "ก่อนสแกน SonarQube และก่อน deploy ทุกรอบ"),
            ("Smoke (หลัง deploy)", "Robot ชุดย่อย", "หน้าเปิดได้ /api/health ตอบ — อ่านอย่างเดียว ไม่สร้างข้อมูล ไม่ส่งเมล", "ทุกครั้งหลังขึ้นระบบจริง"),
        ]),
        ("2. ประเภทและเทคนิคการทดสอบ", [("__table__", ["เทคนิค", "ตัวอย่างในรายงานนี้"]),
            ("Positive (ทางที่ถูก)", "TC_STAFF_LOGIN_001, TC_CUS_LOGIN_001, TC_PROJ_001, TC_BOQ_002"),
            ("Negative (ข้อมูลผิด/ขาด)", "TC_STAFF_LOGIN_002/004/005, TC_CUS_LOGIN_002/004, TC_BOQ_005, TC_PROJ_002"),
            ("Equivalence Partitioning", "เบอร์มือถือ vs เบอร์บ้าน (TC_REG_004) · ชื่อคำเดียว vs ชื่อ-นามสกุล (TC_REG_002) · อีเมลมี/ไม่มีโดเมน (TC_REG_003)"),
            ("Boundary Value", "รหัสผ่านสั้นกว่า 8 ตัว (TC_REG_005) · จำนวน BOQ = 0 (TC_BOQ_006) · ผิดครบ 5 ครั้ง (TC_STAFF_LOGIN_007)"),
            ("Security", "ไม่บอกว่าบัญชีมีจริงไหม (TC_STAFF_LOGIN_003) · ล็อกบัญชี (TC_STAFF_LOGIN_007) · หน้าภายในต้องเข้าสู่ระบบ (TC_*_ACCESS_001) · ลูกค้าเข้าหลังบ้านไม่ได้ (TC_STAFF_ACCESS_002)"),
            ("Calculation / Data", "ยอดบรรทัด 2×(100+50)=300 · กำไร 10% · VAT 7% · ราคาสุทธิ 353.10 (TC_BOQ_001/003)"),
            ("State / Workflow", "เข้าสู่ระบบ → ออกจากระบบ → เปิดหน้าภายในไม่ได้ (TC_STAFF_LOGIN_008, TC_CUS_LOGIN_005) · "
             "ใบเสนอราคา → สัญญา → วางบิล → แจ้งโอน → ใบเสร็จ (TC_SALE_001–009)"),
            ("Cross-browser / Responsive", "Edge (TC_XB_001–003) · มือถือ 390 px (TC_MOB_001–004) · iPad 1024 px (TC_IPAD_001–003)"),
            ("Accessibility", "คีย์บอร์ดล้วน Tab/Enter/Esc · ช่องกรอกและปุ่มมีชื่อ (TC_KEY_C_001–003, TC_KEY_S_001–003)"),
            ("Role-based (สิทธิ์บนหน้าจอ)", "ADMIN ไม่เห็นกำไร/ไอพี/ปุ่มของ OWNER (TC_ROLE_001–004) · ลบถาวรเฉพาะ OWNER (TC_TRASH_004)"),
        ]),
        ("3. การจัดการข้อมูลทดสอบ", [
            ("บัญชีทดสอบ", "robot-tests/seed/seed-robot.ts สร้าง/รีเซ็ตบัญชีทดสอบทั้งหมดและหมวดงาน R / R-01 / R-01-01 ทุกครั้งก่อนรัน (รันซ้ำได้) · "
             "robot_twofa เปิด 2FA ด้วย secret ที่รู้ค่า (Robot คำนวณรหัส 6 หลักเอง) และรหัสสำรอง 2 ชุด"),
            ("ข้อมูลเดินงานขาย/ถังขยะ/Excel", "seed สร้างโครงการร่างพร้อม BOQ ใหม่ทุกรอบ (results/ui-data.json) — "
             "ใบเสนอราคา สัญญา วางบิล ลบ/กู้คืน นำเข้า/ส่งออก ทำผ่านหน้าจอทั้งหมด"),
            ("ข้อมูลที่เทสต์สร้าง", "โครงการใหม่ต่อท้ายชื่อด้วยวันเวลา (เช่น \"โครงการโรบอท 1018202430\") จึงไม่ชนกันข้ามรอบ"),
            ("ข้อมูลส่วนชำระเงิน", "seed สร้างโครงการ \"[โรบอท] ทดสอบชำระเงิน …\" ใหม่ทุกรอบผ่าน API จริง (ใบเสนอราคา → ยอมรับ → สัญญา → ลงนาม → วางบิลงวด 1–3) "
             "และลบของรอบก่อนทิ้ง · ยอดแต่ละงวดเขียนลง results/payment-data.json ให้เทสต์อ่าน ไม่พิมพ์ตัวเลขไว้ในสคริปต์"),
            ("สลิปจำลอง", "ไฟล์ PNG 1×1 จุดที่ชื่อไฟล์บอกผล: mockslip_<ยอด>_<เลขอ้างอิง>_<YYYYMMDD>_<บัญชี>.png · เลขอ้างอิงต่อท้ายด้วยรหัสรอบ กันระบบมองว่าเป็นสลิปซ้ำ"),
            ("ห้ามใช้ข้อมูลจริง", "ไม่ใช้ฐานข้อมูล dev/ระบบจริง · อีเมลทดสอบใช้โดเมน example.test / example.com เท่านั้น"),
        ]),
        ("4. การทดสอบอัตโนมัติ (Automation Approach)", [
            ("โครงสร้าง", "robot-tests/common.robot (ใช้ร่วม) · frontend/ (หน้าบ้าน) · backend/ (หลังบ้าน) · แต่ละโฟลเดอร์มี resource.robot ของตัวเอง"),
            ("ที่มาของเคส", "ทุกเคสเขียน Scenario / Pre / Steps / Data / Expected ไว้ใน [Documentation] — รายงานนี้อ่านจากตรงนั้น จึงตรงกับสคริปต์เสมอ"),
            ("หลักฐาน", "ภาพหน้าจอท้ายทุกเคส (evidence/<รหัสเคส>.png) + ภาพตอนล้ม + log.html ที่บอกทุกขั้นตอน"),
            ("คำสั่งรัน", "node robot-tests/test-env.mjs --run   (seed → เปิดระบบทดสอบ → รันหน้าบ้าน → เปิด API ใหม่ → รันหลังบ้าน → รวมผล → ปิด)\n"
             "python -X utf8 robot-tests/build_report.py --out <ไฟล์.xlsx>   (สร้างรายงานนี้ใหม่)"),
        ]),
        ("5. การจัดการข้อบกพร่อง (Defect Management)", [("__table__", ["ระดับความรุนแรง", "ความหมาย", "ลำดับแก้"]),
            ("Critical / Blocker", "ระบบพัง เข้าใช้งานไม่ได้", "P1 — แก้ทันที ห้าม deploy"),
            ("High", "ฟังก์ชันหลักทำงานไม่ได้ แต่มีทางเลี่ยง", "P1/P2 — แก้ก่อน deploy"),
            ("Medium", "ทำงานผิดพลาดในฟังก์ชันย่อย", "P2 — แก้ในรอบนี้หรือรอบถัดไป"),
            ("Low", "แสดงผลผิดพลาดเล็กน้อย / UI เพี้ยน", "P3 — แก้เมื่อสะดวก"),
        ]),
        ("6. ขั้นตอนเมื่อพบเคสไม่ผ่าน", [
            "ดูภาพหน้าจอและ log.html ก่อน — แยกให้ได้ว่าเป็นบัคของระบบ หรือสคริปต์/ข้อมูลทดสอบผิด",
            "ถ้าเป็นบัคของระบบ: ออก Defect ID (BUG-###) ระบุระดับความรุนแรง ขั้นตอนทำซ้ำ และไฟล์ที่ต้องแก้ ในชีต Defects",
            "ถ้าเป็นของสคริปต์: แก้สคริปต์แล้วรันใหม่ ไม่ออก Defect ID (รอบนี้แก้สคริปต์ไป 5 จุด เช่น รหัสผ่านทดสอบไปซ้ำกับอีเมล)",
            "หลังแก้บัค: รันทั้งชุดใหม่ (ไม่ใช่เฉพาะเคสที่ล้ม) เพื่อยืนยันว่าไม่มีของเดิมพัง",
        ]),
        ("7. ตัวชี้วัด (Metrics)", [
            ("อัตราผ่าน", "จำนวน Pass ÷ จำนวนเคสที่รัน (ชีตสรุปผลคำนวณจากคอลัมน์ Status อัตโนมัติ)"),
            ("ข้อบกพร่องตามระดับ", "นับจากชีต Defects แยก Critical / High / Medium / Low"),
            ("Coverage", "จาก jest --coverage ฝั่ง server (Robot ไม่สร้าง coverage) — ส่งต่อให้ SonarQube"),
        ]),
        ("8. ลำดับเทียบกับ SonarQube และการ deploy", [
            "เขียนเทสต์ → รันเทสต์ + coverage → สแกน SonarQube → แก้ตามผล → รันเทสต์ซ้ำ → สแกนซ้ำ → deploy → smoke test",
            "เทสต์ต้องมาก่อน SonarQube เพราะ SonarQube อ่าน coverage จากผลเทสต์ และการแก้ตาม SonarQube ต้องมีเทสต์คอยยืนยันว่าไม่ทำของเดิมพัง",
        ]),
    ]


# ---------------------------------------------------------------- ชีต Test Cases
TC_COLUMNS = [
    ("Test Case ID", 23), ("Module / Feature", 20), ("Test Scenario / Title", 34), ("Pre-conditions", 30),
    ("Test Steps", 40), ("Test Data", 30), ("Expected Result", 36),
    ("Actual Result", 38), ("Status", 11), ("Defect ID / Bug Link", 14), ("Severity / Priority", 14),
    ("Environment", 26), ("Tester Name", 18), ("Execution Date", 17), ("Screenshot / Evidence Link", image_col_width()), ("Remarks / Note", 36),
]
GROUPS = [
    (1, 7, "1. ข้อมูลระดับโครงสร้างและเคสทดสอบ (Test Setup)", "1E3A8A"),
    (8, 11, "2. ผลการทดสอบและการแจ้งบัค (Execution & Bug Tracking)", "9A3412"),
    (12, 16, "3. การบันทึกและจัดการ (Metadata & Evidence)", "065F46"),
]


def write_test_cases(ws, rows):
    ws.sheet_view.showGridLines = False
    for start, end, label, color in GROUPS:
        ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
        cell = ws.cell(1, start, label)
        cell.font = f(11, True, "FFFFFF")
        cell.fill = fill(color)
        cell.alignment = CENTER
        for c in range(start, end + 1):
            ws.cell(1, c).border = BORDER
            h = ws.cell(2, c, TC_COLUMNS[c - 1][0])
            h.font = f(10, True, "FFFFFF")
            h.fill = fill(color)
            h.alignment = CENTER
            h.border = BORDER
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 30
    for i, (_, width) in enumerate(TC_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    for r, x in enumerate(rows, start=3):
        values = [x["id"], x["module"], x["title"], x["pre"], x["steps"], x["data"], x["expected"],
                  x["actual"], x["status"], x["defect"], x["severity"],
                  MODULE_ENVIRONMENT.get(x["module"], ENVIRONMENT), TESTER, x["date"], "" if x["evidence"] else EVIDENCE_PLACEHOLDER, x["remarks"]]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(r, c, v)
            cell.font = f(9)
            cell.alignment = WRAP_TOP
            cell.border = BORDER
        ws.cell(r, 1).font = f(9, True)
        ws.cell(r, 9).alignment = CENTER
        ws.cell(r, 9).font = f(9, True)
        ws.cell(r, 10).alignment = CENTER
        ws.cell(r, 11).alignment = CENTER
        ws.cell(r, 14).number_format = "yyyy-mm-dd hh:mm:ss"
        if x["defect"]:
            ws.cell(r, 10).hyperlink = Hyperlink(ref=f"J{r}", location=f"'Defects'!A{defect_row(rows, x['defect'])}")
            closed = x["defect_state"] == "Closed"
            ws.cell(r, 10).font = f(9, True, "166534" if closed else "B91C1C")
            if closed:
                ws.cell(r, 10).value = f'{x["defect"]} (ปิดแล้ว)'
        img_h = embed_screenshot(ws, x["evidence"], r, 15)
        # คอลัมน์วันที่ (index 13) เป็น datetime ไม่ห่อบรรทัด จึงไม่นับ
        lines = max(est_lines(v, TC_COLUMNS[c][1] * 1.05) for c, v in enumerate(values) if c != 13)
        # px → pt คูณ 0.75 · แถวต้องสูงพอให้ภาพไม่ล้นไปทับแถวถัดไป
        ws.row_dimensions[r].height = max(min(12.5 * lines + 6, 200), (img_h + CELL_PAD * 2) * 0.75)

    last = len(rows) + 2
    ws.freeze_panes = "B3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(TC_COLUMNS))}{last}"

    status_dv = DataValidation(type="list", formula1='"Pass,Fail,Block,Untested"', allow_blank=False,
                               showErrorMessage=True, errorTitle="Status", error="เลือก Pass / Fail / Block / Untested")
    sev_dv = DataValidation(type="list", allow_blank=True,
                            formula1='"Critical / P1,Blocker / P1,High / P1,High / P2,Medium / P2,Medium / P3,Low / P3"')
    ws.add_data_validation(status_dv)
    ws.add_data_validation(sev_dv)
    status_dv.add(f"I3:I{last + 50}")
    sev_dv.add(f"K3:K{last + 50}")

    rng = f"I3:I{last + 50}"
    for word, bg, fg in [("Pass", "DCFCE7", "166534"), ("Fail", "FEE2E2", "B91C1C"),
                         ("Block", "FEF3C7", "92400E"), ("Untested", "E5E7EB", "374151")]:
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$I3="{word}"'], fill=fill(bg),
                                                       font=Font(name=FONT, bold=True, color=fg)))
    return last


_defect_rows: dict[str, int] = {}


def defect_row(rows, defect_id):
    if not _defect_rows:
        ids = sorted({x["defect"] for x in rows if x["defect"]})
        for i, d in enumerate(ids, start=3):
            _defect_rows[d] = i
    return _defect_rows.get(defect_id, 3)


def write_defects(ws, rows):
    ws.sheet_view.showGridLines = False
    cols = [("Defect ID", 11), ("Test Case ID", 23), ("Title", 34), ("Severity", 10), ("Priority", 9),
            ("Steps to Reproduce", 38), ("Expected Result", 30), ("Actual Result", 34), ("ผลกระทบต่อผู้ใช้", 34),
            ("แนวทางแก้ (Suggested Fix)", 40), ("ไฟล์ที่เกี่ยวข้อง", 30), ("Evidence (ตอนพบบัค)", image_col_width()), ("Found Date", 17),
            ("Defect Status", 12), ("การแก้ไขที่ทำ", 44), ("ผลทดสอบซ้ำ (Retest)", 30), ("หลักฐานหลังแก้", image_col_width())]
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    ws["A1"] = "รายการข้อบกพร่อง (Defect Log) — ออก Defect ID เฉพาะเคสที่วิเคราะห์แล้วว่าเป็นบัคของระบบ"
    ws["A1"].font = f(12, True, NAVY)
    for i, (name, width) in enumerate(cols, start=1):
        c = ws.cell(2, i, name)
        c.font = f(10, True, "FFFFFF")
        c.fill = fill("9A3412")
        c.alignment = CENTER
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[2].height = 30
    grouped: dict[str, list[dict]] = {}
    for x in rows:
        if x["defect"]:
            grouped.setdefault(x["defect"], []).append(x)
    failed = []
    for did in sorted(grouped):
        # เคสที่เขียนมือ (อธิบายผลกระทบละเอียด) ขึ้นเป็นตัวอย่างก่อนเคสสิทธิ์ที่สร้างอัตโนมัติ
        grouped[did].sort(key=lambda g: ("_ACL_" in g["id"], g["id"]))
        first = dict(grouped[did][0])
        if len(grouped[did]) > 1:
            first["id"] = "\n".join(g["id"] for g in grouped[did])
        failed.append(first)
    for r, x in enumerate(failed, start=3):
        d = x["defect_info"]
        sev = d["severity"]
        closed = x["defect_state"] == "Closed"
        # ปิดแล้ว: Actual/หลักฐาน/วันที่พบ = ของรอบแรกที่พบบัค · ผลรอบนี้ไปอยู่ช่อง Retest
        found_actual = x["before_actual"] if closed else x["actual"]
        found_evidence = x["before_evidence"] if closed else x["evidence"]
        found_date = x["before_date"] if closed else x["date"]
        retest = (f'Pass ({x["date"]:%Y-%m-%d %H:%M}) — {x["actual"]}' if closed else "ยังไม่ได้แก้")
        values = [d["id"], x["id"], d["title"], sev, d["priority"], x["steps"], x["expected"], found_actual,
                  d["impact"], d["fix"], d["where"], "" if found_evidence else "ไม่มีภาพ", found_date,
                  x["defect_state"], d.get("change", "") if closed else "", retest, ""]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(r, c, v)
            cell.font = f(9, bold=(c == 1))
            cell.alignment = WRAP_TOP
            cell.border = BORDER
        ws.cell(r, 13).number_format = "yyyy-mm-dd hh:mm"
        ws.cell(r, 14).alignment = CENTER
        ws.cell(r, 14).font = f(9, True, "166534" if closed else "B91C1C")
        ws.cell(r, 14).fill = fill("DCFCE7" if closed else "FEE2E2")
        h_before = embed_screenshot(ws, found_evidence, r, 12)
        h_after = embed_screenshot(ws, x["evidence"], r, 17) if closed else 0
        ws.row_dimensions[r].height = max(135, (max(h_before, h_after) + CELL_PAD * 2) * 0.75)
    dv = DataValidation(type="list", formula1='"Open,In Progress,Fixed,Retest,Closed,Won\'t Fix"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("N3:N60")
    sev_dv = DataValidation(type="list", formula1='"Critical,Blocker,High,Medium,Low"', allow_blank=True)
    ws.add_data_validation(sev_dv)
    sev_dv.add("D3:D60")
    ws.freeze_panes = "C3"
    return len(failed)


# ---------------------------------------------------------------- ชีตสรุปผล (สูตรทั้งหมด อ่านจากชีต Test Cases/Defects)
def write_summary(ws, rows, meta, tc_last):
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGH", [3, 34, 14, 14, 14, 14, 14, 16]):
        ws.column_dimensions[col].width = w
    ws["B2"] = "รายงานผลการทดสอบระบบ SystemBOQ (Automated Testing)"
    ws["B2"].font = f(16, True, NAVY)
    ws["B3"] = "ระบบประเมินราคาต้นทุนงานก่อสร้าง — หน้าบ้าน (เว็บลูกค้า) และหลังบ้าน (ระบบพนักงาน)"
    ws["B3"].font = f(10, color="6B7280")

    info = [
        ("รอบการทดสอบ", f'{meta["start"].strftime("%Y-%m-%d %H:%M")} – {meta["end"].strftime("%H:%M")} '
                        f'(ใช้เวลา {int(meta["elapsed"].total_seconds())} วินาที)'),
        ("เครื่องมือ", "Robot Framework 7.5 + SeleniumLibrary 6.9 (Python 3.13)"),
        ("สภาพแวดล้อม", ENVIRONMENT_FULL),
        ("ผู้รับผิดชอบการทดสอบ", ""),
        ("รายงานละเอียดของ Robot", "testcase-evidence/report.html"),
    ]
    r = 5
    for k, v in info:
        ws.cell(r, 2, k).font = f(10, True)
        cell = ws.cell(r, 3, v)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)
        cell.font = f(10)
        if k == "ผู้รับผิดชอบการทดสอบ":
            for c in range(3, 9):
                ws.cell(r, c).fill = fill("FFFF00")
            cell.value = None
        if k == "รายงานละเอียดของ Robot":
            cell.hyperlink = v
            cell.font = Font(name=FONT, size=10, color="1D4ED8", underline="single")
        r += 1
    ws.cell(r, 2, "ช่องสีเหลือง = ให้กรอกเอง · ตัวเลขด้านล่างเป็นสูตรที่นับจากชีต Test Cases (แก้คอลัมน์ Status แล้วอัปเดตเอง)").font = f(9, italic=True, color="6B7280")
    r += 2

    rng_status = f"'Test Cases'!$I$3:$I${tc_last + 50}"
    rng_module = f"'Test Cases'!$B$3:$B${tc_last + 50}"
    rng_id = f"'Test Cases'!$A$3:$A${tc_last + 50}"

    ws.cell(r, 2, "ภาพรวม").font = f(12, True, NAVY)
    r += 1
    heads = ["", "ทั้งหมด", "Pass", "Fail", "Block", "Untested", "อัตราผ่าน"]
    for i, h in enumerate(heads):
        c = ws.cell(r, 2 + i, h or "ตัวชี้วัด")
        c.font = f(10, True, "FFFFFF")
        c.fill = fill(NAVY)
        c.alignment = CENTER
        c.border = BORDER
    r += 1
    ws.cell(r, 2, "จำนวนเคส").font = f(10, True)
    ws.cell(r, 3, f'=COUNTA({rng_id})')
    for i, word in enumerate(["Pass", "Fail", "Block", "Untested"]):
        ws.cell(r, 4 + i, f'=COUNTIF({rng_status},"{word}")')
    ws.cell(r, 8, f"=IF(C{r}-G{r}=0,0,D{r}/(C{r}-G{r}))")
    ws.cell(r, 8).number_format = "0.0%"
    for c in range(2, 9):
        ws.cell(r, c).border = BORDER
        ws.cell(r, c).alignment = CENTER if c > 2 else WRAP_TOP
        if c > 2:
            ws.cell(r, c).font = f(12, True)
    overview_row = r
    r += 1
    ws.cell(r, 2, "อัตราผ่าน = Pass ÷ (ทั้งหมด − Untested)").font = f(8, italic=True, color="6B7280")
    r += 2

    history = meta.get("history") or []
    if history:
        ws.cell(r, 2, "ประวัติการรัน").font = f(12, True, NAVY)
        r += 1
        for i, h in enumerate(["รอบ", "วันเวลา", "ทั้งหมด", "Pass", "Fail", "Block", "อัตราผ่าน"]):
            c = ws.cell(r, 2 + i, h)
            c.font = f(10, True, "FFFFFF")
            c.fill = fill(NAVY)
            c.alignment = CENTER
            c.border = BORDER
        first_hist = r + 1
        # รอบก่อน ๆ = ค่าจาก testcase-evidence-round<N>/output.xml (ผลรันจริงที่เก็บไว้ ไม่ได้พิมพ์เอง)
        for h in history:
            r += 1
            ws.cell(r, 2, f'รอบที่ {h["n"]} — {h["label"]}')
            ws.cell(r, 3, h["start"].strftime("%Y-%m-%d %H:%M"))
            for i, k in enumerate(["total", "pass", "fail", "skip"]):
                ws.cell(r, 4 + i, h[k])
            ws.cell(r, 8, f"=IF(D{r}=0,0,E{r}/D{r})")
            ws.cell(r, 2).comment = Comment(f'ค่าของรอบนี้อ่านจาก testcase-evidence-round{h["n"]}/output.xml (ผลรัน Robot จริง)',
                                            "build_report.py")
        r += 1
        ws.cell(r, 2, f'รอบที่ {len(history) + 1} — {meta["label"]}')
        ws.cell(r, 3, meta["start"].strftime("%Y-%m-%d %H:%M"))
        for i, col in enumerate(["C", "D", "E", "F"]):
            ws.cell(r, 4 + i, f"={col}{overview_row}")
        ws.cell(r, 8, f"=IF(D{r}=0,0,E{r}/D{r})")
        for rr in range(first_hist, r + 1):
            ws.cell(rr, 8).number_format = "0.0%"
            for c in range(2, 9):
                ws.cell(rr, c).border = BORDER
                ws.cell(rr, c).font = f(10)
                if c > 2:
                    ws.cell(rr, c).alignment = CENTER
        r += 2

    ws.cell(r, 2, "แยกตามส่วนของระบบ / โมดูล").font = f(12, True, NAVY)
    r += 1
    for i, h in enumerate(["Module / Feature", "ทั้งหมด", "Pass", "Fail", "Block", "Untested", "อัตราผ่าน"]):
        c = ws.cell(r, 2 + i, h)
        c.font = f(10, True, "FFFFFF")
        c.fill = fill(NAVY)
        c.alignment = CENTER
        c.border = BORDER
    r += 1
    modules: list[str] = []
    for x in rows:
        if x["module"] not in modules:
            modules.append(x["module"])
    first_mod = r
    for m in modules:
        ws.cell(r, 2, m).font = f(10)
        ws.cell(r, 3, f'=COUNTIF({rng_module},B{r})')
        for i, word in enumerate(["Pass", "Fail", "Block", "Untested"]):
            ws.cell(r, 4 + i, f'=COUNTIFS({rng_module},$B{r},{rng_status},"{word}")')
        ws.cell(r, 8, f"=IF(C{r}-G{r}=0,0,D{r}/(C{r}-G{r}))")
        ws.cell(r, 8).number_format = "0.0%"
        for c in range(2, 9):
            ws.cell(r, c).border = BORDER
            if c > 2:
                ws.cell(r, c).alignment = CENTER
                ws.cell(r, c).font = f(10)
        r += 1
    ws.cell(r, 2, "รวม").font = f(10, True)
    for c in range(3, 8):
        col = get_column_letter(c)
        ws.cell(r, c, f"=SUM({col}{first_mod}:{col}{r - 1})")
    ws.cell(r, 8, f"=IF(C{r}-G{r}=0,0,D{r}/(C{r}-G{r}))")
    ws.cell(r, 8).number_format = "0.0%"
    for c in range(2, 9):
        ws.cell(r, c).border = BORDER
        ws.cell(r, c).fill = fill("EEF2FF")
        if c > 2:
            ws.cell(r, c).alignment = CENTER
            ws.cell(r, c).font = f(10, True)
    ws.cell(r + 1, 2, f'=IF(C{r}=C{overview_row},"✓ ผลรวมรายโมดูลตรงกับภาพรวม","✗ ผลรวมรายโมดูลไม่ตรงกับภาพรวม — ตรวจคอลัมน์ Module")').font = f(8, italic=True, color="6B7280")
    r += 3

    ws.cell(r, 2, "ข้อบกพร่องที่พบ").font = f(12, True, NAVY)
    r += 1
    for i, h in enumerate(["ระดับความรุนแรง", "จำนวน", "ยังไม่ปิด"]):
        c = ws.cell(r, 2 + i, h)
        c.font = f(10, True, "FFFFFF")
        c.fill = fill("9A3412")
        c.alignment = CENTER
        c.border = BORDER
    r += 1
    first_sev = r
    for sev in ["Critical", "Blocker", "High", "Medium", "Low"]:
        ws.cell(r, 2, sev).font = f(10)
        ws.cell(r, 3, f'=COUNTIF(Defects!$D$3:$D$60,B{r})')
        ws.cell(r, 4, f'=COUNTIFS(Defects!$D$3:$D$60,B{r},Defects!$N$3:$N$60,"<>Closed")')
        for c in range(2, 5):
            ws.cell(r, c).border = BORDER
            if c > 2:
                ws.cell(r, c).alignment = CENTER
        r += 1
    ws.cell(r, 2, "รวม").font = f(10, True)
    ws.cell(r, 3, f"=SUM(C{first_sev}:C{r - 1})")
    ws.cell(r, 4, f"=SUM(D{first_sev}:D{r - 1})")
    for c in range(2, 5):
        ws.cell(r, c).border = BORDER
        ws.cell(r, c).fill = fill("FFF7ED")
        if c > 2:
            ws.cell(r, c).alignment = CENTER
            ws.cell(r, c).font = f(10, True)
    total_sev_row = r
    r += 2

    ws.cell(r, 2, "ผ่านเกณฑ์จบการทดสอบหรือยัง").font = f(12, True, NAVY)
    r += 1
    crit_high = f"SUMPRODUCT((Defects!$D$3:$D$60={{\"Critical\",\"Blocker\",\"High\"}})*(Defects!$N$3:$N$60<>\"Closed\"))"
    checks = [
        ("รันครบทุกเคส (ไม่มี Untested)", f'=IF(G{overview_row}=0,"ผ่าน","ยังไม่ผ่าน")'),
        ("ไม่มีข้อบกพร่อง Critical/Blocker/High ค้าง", f'=IF({crit_high}=0,"ผ่าน","ยังไม่ผ่าน")'),
        ("เคสที่ไม่ผ่านทุกเคสมี Defect ID",
         f"=IF(COUNTIFS({rng_status},\"Fail\",'Test Cases'!$J$3:$J${tc_last + 50},\"\")=0,\"ผ่าน\",\"ยังไม่ผ่าน\")"),
        ("ข้อบกพร่องทุกตัวปิดแล้ว (พร้อม deploy)", f'=IF(D{total_sev_row}=0,"ผ่าน","ยังไม่ผ่าน")'),
    ]
    first_check = r
    for label, formula in checks:
        ws.cell(r, 2, label).font = f(10)
        ws.cell(r, 3, formula)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
        ws.cell(r, 3).alignment = CENTER
        for c in range(2, 5):
            ws.cell(r, c).border = BORDER
        r += 1
    ws.conditional_formatting.add(f"C{first_check}:C{r - 1}", FormulaRule(
        formula=[f'C{first_check}="ผ่าน"'], fill=fill("DCFCE7"), font=Font(name=FONT, bold=True, color="166534")))
    ws.conditional_formatting.add(f"C{first_check}:C{r - 1}", FormulaRule(
        formula=[f'C{first_check}="ยังไม่ผ่าน"'], fill=fill("FEE2E2"), font=Font(name=FONT, bold=True, color="B91C1C")))


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True)
    ap.add_argument("--round1", help="โฟลเดอร์ผลรันรอบก่อน (ค่าเริ่มต้น: testcase-evidence-round1 ข้างไฟล์ Excel)")
    ap.add_argument("--label", default="รอบล่าสุด", help="คำอธิบายรอบที่กำลังรายงาน เช่น \"หลังปรับส่วนชำระเงิน\"")
    ap.add_argument("--note", default="", help="หมายเหตุของรอบนี้ (ขึ้นในแผนการทดสอบ หัวข้อขอบเขต)")
    args = ap.parse_args()
    out = Path(args.out)
    global BASE_DIR
    BASE_DIR = out.parent
    round1_dir = Path(args.round1) if args.round1 else out.parent / "testcase-evidence-round1"
    out.parent.mkdir(parents=True, exist_ok=True)

    rows, meta = collect(round1_dir)

    # ประวัติ: ทุกโฟลเดอร์ testcase-evidence-round<N> ข้างไฟล์ Excel เรียงตามเลขรอบ
    history = []
    for d in out.parent.glob("testcase-evidence-round*"):
        m = re.fullmatch(r"testcase-evidence-round(\d+)", d.name)
        if not m:
            continue
        _, stats = read_round(d / "output.xml", UI_SUITES)  # รอบหลัง ๆ มีชุด API ในไฟล์เดียวกัน — นับเฉพาะหน้าจอ
        if stats:
            n = int(m.group(1))
            history.append({"n": n, "label": ROUND_LABELS.get(n, "รอบก่อนหน้า"), **stats})
    meta["history"] = sorted(history, key=lambda h: h["n"])
    meta["label"] = args.label
    meta["note"] = args.note

    # หลักฐานไปไว้ข้างไฟล์ Excel (ลิงก์ในไฟล์เป็นแบบสัมพัทธ์)
    ev = out.parent / "testcase-evidence"
    if ev.exists():
        shutil.rmtree(ev)
    shutil.copytree(RESULTS, ev, ignore=shutil.ignore_patterns("env-*.log"))

    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "สรุปผล"
    ws_plan = wb.create_sheet("แผนการทดสอบ")
    ws_strat = wb.create_sheet("กลยุทธ์การทดสอบ")
    ws_tc = wb.create_sheet("Test Cases")
    ws_def = wb.create_sheet("Defects")

    write_doc_sheet(ws_plan, "แผนการทดสอบ (Test Plan)", "SystemBOQ — การทดสอบอัตโนมัติก่อนสแกน SonarQube และก่อนขึ้นระบบจริง",
                    plan_sections(rows, meta))
    write_doc_sheet(ws_strat, "กลยุทธ์การทดสอบ (Test Strategy)", "ทดสอบอะไร ด้วยวิธีไหน ระดับไหน และจัดการผลอย่างไร",
                    strategy_sections())
    tc_last = write_test_cases(ws_tc, rows)
    write_defects(ws_def, rows)
    write_summary(ws_sum, rows, meta, tc_last)

    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = {"สรุปผล": NAVY, "Test Cases": "1E3A8A", "Defects": "9A3412"}.get(ws.title, "6B7280")
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws_tc.print_title_rows = "1:2"

    wb.save(out)
    passed = sum(1 for x in rows if x["status"] == "Pass")
    print(f"wrote {out} · {len(rows)} cases · pass {passed} · fail {sum(1 for x in rows if x['status'] == 'Fail')}")


if __name__ == "__main__":
    main()
