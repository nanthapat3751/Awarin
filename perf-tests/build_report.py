"""
สร้างรายงาน Performance Test (Excel) จากผล k6 ใน perf-tests/run-all.sh

    python perf-tests/build_report.py <โฟลเดอร์ผล> <ไฟล์ปลายทาง.xlsx> [findings.json]
        [--compare <โฟลเดอร์ผลรอบก่อน>] [--first-pdf <first-pdf.csv>] [--round <ชื่อรอบ>]

--compare เพิ่มแผ่นงาน "เทียบรอบ" (รอบก่อน vs รอบนี้) · --first-pdf = ผลวัด PDF ใบแรกหลังเปิด API (first-pdf.sh)

อ่านต่อรูปแบบ (smoke/load/stress/spike/pdf):
    <type>-summary.json    ผลสรุปจาก k6 handleSummary
    <type>-requests.csv    ทุกคำขอ (k6 --out csv) — ใช้คิดเปอร์เซ็นไทล์รายเส้นและรายช่วงเวลา
    <type>-resources.csv   CPU/RAM ของ API + ฐานข้อมูลทุก 5 วินาที (sample-resources.ps1)

เปอร์เซ็นไทล์ (p50/p90/p95/p99) คิดจากข้อมูลดิบทุกคำขอใน Python เพราะ Excel ไม่มีข้อมูลดิบ
ช่องที่เป็นผลตัดสิน/อัตราส่วน/ผลรวมเป็นสูตร Excel อ้างค่าที่อยู่ในแผ่นงาน
"""
from __future__ import annotations

import csv
import gzip
import json
import math
import re
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Tahoma"  # มีตัวอักษรไทยในตัว (เหมือนรายงาน Robot เดิม)
NAVY = "000099"
TYPES = ["smoke", "load", "stress", "spike", "pdf"]
TYPE_TH = {
    "smoke": "Smoke test — ตรวจว่าทุกเส้นทำงาน",
    "load": "Load test — ใช้งานปกติต่อเนื่อง",
    "stress": "Stress test — เพิ่มผู้ใช้จนเห็นขีดจำกัด",
    "spike": "Spike test — ผู้ใช้พุ่งขึ้นกะทันหัน",
    "pdf": "PDF test — สร้าง PDF ใบเสนอราคาพร้อมกัน",
}
# ผู้ใช้สูงสุดต่อกลุ่ม (staff และ visitor ไต่พร้อมกันเท่ากัน — pdf มีกลุ่มเดียว)
PEAK_PER_GROUP = {"smoke": 1, "load": 20, "stress": 150, "spike": 150, "pdf": 5}
GROUPS = {"smoke": 2, "load": 2, "stress": 2, "spike": 2, "pdf": 1}
ENDPOINT_TH = {
    "public:settings": "ค่าตั้งเว็บ (ชื่อบริษัท รูปหน้าเว็บ)",
    "public:house-designs-featured": "แบบบ้านแนะนำ (หน้าแรก)",
    "public:house-designs": "รายการแบบบ้าน",
    "public:portfolio": "รายการผลงาน",
    "staff:me": "ข้อมูลผู้ใช้ที่ล็อกอิน",
    "staff:dashboard-summary": "แดชบอร์ด — สรุป",
    "staff:dashboard-pipeline": "แดชบอร์ด — ท่อขาย",
    "staff:notifications": "การแจ้งเตือน",
    "staff:projects-list": "รายการโครงการ",
    "staff:project-detail": "รายละเอียดโครงการ",
    "staff:project-boq": "BOQ ของโครงการ (178 รายการ)",
    "staff:project-quotes": "ใบเสนอราคาของโครงการ",
    "staff:materials": "คลังวัสดุ (121 รายการ)",
    "staff:customers": "แฟ้มลูกค้า",
    "staff:leads": "คำขอลูกค้า",
    "pdf:quote": "สร้าง PDF ใบเสนอราคา",
}

thin = Side(style="thin", color="CBD5E1")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def font(size=10, bold=False, color="1F2937", italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


def header(ws, row, values, widths=None):
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = font(10, True, "FFFFFF")
        c.fill = fill(NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 32


def body(c, fmt=None, bold=False, align="left", color="1F2937"):
    c.font = font(10, bold, color)
    c.border = BORDER
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    if fmt:
        c.number_format = fmt


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = font(14, True, NAVY)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = font(10, False, "6B7280", italic=True)


def show_axes(*charts):
    """openpyxl 3.1 ตั้งแกนเป็น delete=True — ไม่เปิดเองกราฟจะไม่มีตัวเลขบนแกน"""
    for ch in charts:
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        ch.y_axis.numFmt = "#,##0"
        # เส้นตรงจุดต่อจุด — เส้นโค้งแบบ smooth ทำให้กราฟดูมีค่าที่ไม่ได้วัดจริง
        for series in ch.series:
            series.smooth = False


def thai_date(dt: datetime) -> str:
    return f"{dt:%d/%m}/{dt.year + 543} {dt:%H:%M}"


def pct(values, p):
    if not values:
        return None
    s = sorted(values)
    k = (len(s) - 1) * p / 100
    lo, hi = math.floor(k), math.ceil(k)
    return s[lo] if lo == hi else s[lo] + (s[hi] - s[lo]) * (k - lo)


def read_requests(path: Path):
    """คืน (คำขอทั้งหมด, vus ต่อวินาที) — ตัดคำขอ setup:login ออก"""
    reqs = []  # (ts, name, duration_ms, failed, status)
    durations = {}
    failed = {}
    vus = {}
    # ผลที่เก็บเป็นหลักฐาน (results-round*/) บีบเป็น .csv.gz
    opener = gzip.open(path, "rt", encoding="utf-8") if path.suffix == ".gz" else path.open(encoding="utf-8")
    with opener as f:
        for r in csv.DictReader(f):
            name = r["name"]
            m = r["metric_name"]
            ts = int(float(r["timestamp"]))
            if m == "vus":
                vus[ts] = max(vus.get(ts, 0), float(r["metric_value"]))
                continue
            if name.startswith("setup:") or not name:
                continue
            key = (ts, name, r["url"], r["status"])
            if m == "http_req_duration":
                durations.setdefault(key, []).append(float(r["metric_value"]))
            elif m == "http_req_failed":
                failed.setdefault(key, []).append(float(r["metric_value"]))
    for key, ds in durations.items():
        fs = failed.get(key, [0.0] * len(ds))
        for i, d in enumerate(ds):
            reqs.append((key[0], key[1], d, fs[i] if i < len(fs) else 0.0, key[3]))
    return reqs, vus


def read_resources(path: Path):
    rows = []
    if not path.exists():
        return rows
    with path.open(encoding="ascii", errors="ignore") as f:
        for r in csv.DictReader(f):
            try:
                # PowerShell "o" ให้เศษวินาที 7 หลัก — ตัดเหลือ 6 ให้ Python อ่านได้ทุกรุ่น
                raw = re.sub(r"(\.\d{6})\d+", r"\1", r["time"])
                rows.append(
                    (
                        datetime.fromisoformat(raw),
                        float(r["api_cpu_pct"] or 0),
                        float(r["api_mem_mb"] or 0),
                        float(r["db_cpu_pct"] or 0),
                        float(r["db_mem_mb"] or 0),
                    )
                )
            except ValueError:
                continue
    return rows


def load_round(src: Path) -> dict:
    data = {}
    for t in TYPES:
        sp = src / f"{t}-summary.json"
        rp = src / f"{t}-requests.csv"
        if not rp.exists():
            rp = src / f"{t}-requests.csv.gz"
        if not sp.exists() or not rp.exists():
            continue
        summary = json.loads(sp.read_text(encoding="utf-8"))
        reqs, vus = read_requests(rp)
        data[t] = {
            "summary": summary,
            "reqs": reqs,
            "vus": vus,
            "res": read_resources(src / f"{t}-resources.csv"),
        }
    return data


def round_stats(d: dict) -> dict:
    """ตัวเลขหลักต่อรูปแบบ — p95 ทั้งหมด / p95 เฉพาะเส้นพนักงาน / ข้อมูลที่รับ (MB)"""
    out = {}
    for t, v in d.items():
        m = v["summary"]["metrics"]
        all_d = [x[2] for x in v["reqs"]]
        staff = [x[2] for x in v["reqs"] if x[1].startswith("staff:")]
        out[t] = {
            "rps": m["http_reqs"]["values"]["rate"],
            "p95": pct(all_d, 95),
            "p95_staff": pct(staff, 95) if staff else None,
            "max": max(all_d) if all_d else None,
            "mb": m["data_received"]["values"]["count"] / 1_000_000,
            # ต่อคำขอ = เทียบข้ามรอบได้ตรงกว่ายอดรวม (รอบที่เร็วขึ้นยิงได้มากคำขอขึ้น ยอดรวมจึงเพิ่มได้)
            "kb_per_req": m["data_received"]["values"]["count"] / 1000 / max(1, m["http_reqs"]["values"]["count"]),
            "err": m["http_req_failed"]["values"]["rate"],
        }
    return out


def endpoint_p95(d: dict, t: str) -> dict:
    by = defaultdict(list)
    for _ts, name, dur, _fl, _st in d.get(t, {}).get("reqs", []):
        by[name].append(dur)
    return {k: pct(v, 95) for k, v in by.items()}


def change_colors(ws, ref: str, first_cell: str, good_when: str):
    """ฟอนต์เขียว = ดีขึ้น · แดง = แย่ลง · good_when = "<0" (เวลา) หรือ ">0" (คำขอ/วินาที)"""
    bad_when = ">0" if good_when == "<0" else "<0"
    for cond, color in ((good_when, "166534"), (bad_when, "B91C1C")):
        ws.conditional_formatting.add(
            ref, FormulaRule(formula=[f"AND(ISNUMBER({first_cell}),{first_cell}{cond})"], font=Font(name=FONT, bold=True, color=color))
        )


def compare_sheet(wb, prev: dict, cur: dict, first_pdf: Path | None, label_prev: str, label_cur: str):
    ws = wb.create_sheet("เทียบรอบ", 1)
    ws.sheet_properties.tabColor = "B45309"
    ws.sheet_view.showGridLines = False
    title(ws, f"เทียบผล {label_prev} กับ {label_cur}",
          "เปลี่ยน % = (รอบนี้ − รอบก่อน) ÷ รอบก่อน · ฟอนต์เขียว = ดีขึ้น · แดง = แย่ลง · ข้อมูลรวมขึ้นกับจำนวนคำขอ — ดูแถว 'ข้อมูลต่อคำขอ' ประกอบ")
    header(ws, 4, ["รูปแบบ", "ตัวชี้วัด", label_prev, label_cur, "เปลี่ยน %"], [38, 30, 14, 14, 12])
    sp, sc = round_stats(prev), round_stats(cur)
    metrics = [
        ("rps", "คำขอ/วินาที", "#,##0.0"),
        ("p95", "p95 ทุกคำขอ (ms)", "#,##0.0"),
        ("p95_staff", "p95 เส้นพนักงาน (ms)", "#,##0.0"),
        ("max", "ช้าสุด (ms)", "#,##0.0"),
        ("mb", "ข้อมูลที่ส่งกลับรวม (MB)", "#,##0.0"),
        ("kb_per_req", "ข้อมูลต่อคำขอ (KB)", "#,##0.0"),
        ("err", "อัตรา error", "0.00%"),
    ]
    r = 4
    p95_rows = {}
    for t in TYPES:
        if t not in sp or t not in sc:
            continue
        for key, label, fmt in metrics:
            a, b = sp[t][key], sc[t][key]
            if a is None and b is None:
                continue
            r += 1
            body(ws.cell(row=r, column=1, value=TYPE_TH[t]), bold=key == "rps")
            body(ws.cell(row=r, column=2, value=label))
            body(ws.cell(row=r, column=3, value=a), fmt, align="right")
            body(ws.cell(row=r, column=4, value=b), fmt, align="right")
            body(ws.cell(row=r, column=5, value=f'=IF(OR(C{r}="",C{r}=0),"-",(D{r}-C{r})/C{r})'), "+0.0%;-0.0%;0.0%", align="right")
            change_colors(ws, f"E{r}", f"E{r}", ">0" if key == "rps" else "<0")
            if key == "p95":
                p95_rows[t] = r

    # p95 รายเส้น ของ load และ stress
    r += 2
    ws.cell(row=r, column=1, value="p95 รายเส้น API (ms)").font = font(11, True, NAVY)
    r += 1
    header(ws, r, ["รูปแบบ · เส้น API", "ทำอะไร", label_prev, label_cur, "เปลี่ยน %"])
    ep_start = r + 1
    for t in ("load", "stress"):
        a, b = endpoint_p95(prev, t), endpoint_p95(cur, t)
        for name in sorted(set(a) & set(b), key=lambda n: (n.split(":")[0] != "public", n)):
            r += 1
            body(ws.cell(row=r, column=1, value=f"{t} · {name}"))
            body(ws.cell(row=r, column=2, value=ENDPOINT_TH.get(name, "")))
            body(ws.cell(row=r, column=3, value=a[name]), "#,##0.0", align="right")
            body(ws.cell(row=r, column=4, value=b[name]), "#,##0.0", align="right")
            body(ws.cell(row=r, column=5, value=f'=IF(C{r}=0,"-",(D{r}-C{r})/C{r})'), "+0.0%;-0.0%;0.0%", align="right")
    if r >= ep_start:
        change_colors(ws, f"E{ep_start}:E{r}", f"E{ep_start}", "<0")

    # PDF ใบแรกหลังเปิด API ใหม่
    if first_pdf and first_pdf.exists():
        r += 2
        ws.cell(row=r, column=1, value="PDF ใบแรกหลังเปิด API ใหม่ — อุ่น Chromium ตอนเปิด (PDF_WARMUP) ปิด vs เปิด").font = font(11, True, NAVY)
        r += 1
        header(ws, r, ["อุ่น Chromium", "ครั้งที่เปิด API", "PDF ใบที่", "วินาที"])
        fp_start = r + 1
        with first_pdf.open(encoding="utf-8") as f:
            for row in csv.DictReader(f):
                r += 1
                body(ws.cell(row=r, column=1, value="เปิด" if row["warmup"] == "on" else "ปิด"))
                body(ws.cell(row=r, column=2, value=int(row["attempt"])), "0", align="right")
                body(ws.cell(row=r, column=3, value=int(row["pdf_no"])), "0", align="right")
                body(ws.cell(row=r, column=4, value=float(row["seconds"])), "0.00", align="right")
        fp_end = r
        for lab, mode in (("ใบแรก เฉลี่ย — ปิด", "ปิด"), ("ใบแรก เฉลี่ย — เปิด", "เปิด")):
            r += 1
            body(ws.cell(row=r, column=1, value=lab), bold=True)
            body(ws.cell(row=r, column=4, value=f'=AVERAGEIFS(D{fp_start}:D{fp_end},A{fp_start}:A{fp_end},"{mode}",C{fp_start}:C{fp_end},1)'), "0.00", bold=True, align="right")
        r += 1
        note = ws.cell(row=r, column=1, value="ครั้งที่ 1 = เปิดเครื่องครั้งแรก (ไฟล์ของ Chromium ยังไม่อยู่ในแคชของระบบ) ใกล้เคียงตอน deploy คอนเทนเนอร์ใหม่บน VPS · ครั้งที่ 2–4 แคชอุ่นแล้ว")
        note.font = font(9, italic=True, color="6B7280")

    # กราฟ p95 รายรูปแบบ — ตารางเล็กด้านขวาอ้างค่าจากตารางหลัก
    if p95_rows:
        rows = [p95_rows[t] for t in TYPES if t in p95_rows]
        for col, v in ((8, "รูปแบบ"), (9, label_prev), (10, label_cur)):
            ws.cell(row=4, column=col, value=v).font = font(9, True)
            ws.column_dimensions[get_column_letter(col)].width = 14
        for i, row in enumerate(rows, start=5):
            ws.cell(row=i, column=8, value=f'=LEFT(A{row},FIND(" ",A{row})-1)')
            ws.cell(row=i, column=9, value=f"=C{row}").number_format = "#,##0"
            ws.cell(row=i, column=10, value=f"=D{row}").number_format = "#,##0"
        last_c = 4 + len(rows)
        ch = BarChart()
        ch.type = "col"
        ch.grouping = "clustered"
        ch.title = "p95 ทุกคำขอ (ms) — รอบก่อน vs รอบนี้"
        ch.height, ch.width = 8.5, 20
        ch.add_data(Reference(ws, min_col=9, max_col=10, min_row=4, max_row=last_c), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=8, min_row=5, max_row=last_c))
        ch.y_axis.title = "ms"
        ch.legend.position = "b"
        show_axes(ch)
        ws.add_chart(ch, "H12")


def main():
    args = sys.argv[1:]
    opts = {}
    for flag in ("--compare", "--first-pdf", "--round"):
        if flag in args:
            i = args.index(flag)
            opts[flag] = args[i + 1]
            del args[i : i + 2]
    src = Path(args[0])
    out = Path(args[1])
    findings = json.loads(Path(args[2]).read_text(encoding="utf-8")) if len(args) > 2 else {}
    round_label = opts.get("--round", "")

    data = load_round(src)

    wb = Workbook()
    ws = wb.active
    ws.title = "สรุปผล"
    ws_plan = wb.create_sheet("วิธีทดสอบและเกณฑ์")
    ws_ep = wb.create_sheet("ผลรายเส้น API")
    ws_st = wb.create_sheet("ผลตามช่วงเวลา")
    ws_res = wb.create_sheet("ทรัพยากรเครื่อง")
    ws_fd = wb.create_sheet("ข้อสังเกต")
    for s, color in zip(wb.worksheets, ["000099", "374151", "065F46", "1D4ED8", "9A3412", "6B21A8"]):
        s.sheet_properties.tabColor = color
        s.sheet_view.showGridLines = False

    # ---------- วิธีทดสอบและเกณฑ์ ----------
    title(ws_plan, "วิธีทดสอบและเกณฑ์ตัดสิน", "Performance test ระบบ SystemBOQ บนเครื่องพัฒนา (ไม่ได้ทดสอบเว็บจริงที่ deploy)")
    env = findings.get("environment", [])
    r = 4
    ws_plan.cell(row=r, column=1, value="สภาพแวดล้อม").font = font(11, True, NAVY)
    r += 1
    header(ws_plan, r, ["หัวข้อ", "รายละเอียด"], [30, 100])
    for k, v in env:
        r += 1
        body(ws_plan.cell(row=r, column=1, value=k), bold=True)
        body(ws_plan.cell(row=r, column=2, value=v))
    r += 2
    ws_plan.cell(row=r, column=1, value="เกณฑ์ตัดสิน (แก้ตัวเลขได้ — แผ่นงานอื่นอ้างช่องเหล่านี้)").font = font(11, True, NAVY)
    r += 1
    header(ws_plan, r, ["เกณฑ์", "ค่า"])
    crit_rows = {}
    for key, label, value, fmt in [
        ("read_p95", "เวลาตอบ p95 ของเส้นอ่านข้อมูล ไม่เกิน (ms)", 800, "#,##0"),
        ("pdf_p95", "เวลาตอบ p95 ของการสร้าง PDF ไม่เกิน (ms)", 15000, "#,##0"),
        ("err", "อัตรา error ไม่เกิน", 0.01, "0.0%"),
    ]:
        r += 1
        body(ws_plan.cell(row=r, column=1, value=label))
        c = ws_plan.cell(row=r, column=2, value=value)
        body(c, fmt, color="0000FF")
        c.fill = fill("FEF9C3")
        crit_rows[key] = f"'วิธีทดสอบและเกณฑ์'!$B${r}"
    r += 1
    note = ws_plan.cell(row=r, column=1, value="ที่มาของเกณฑ์: ผู้ทดสอบกำหนดเอง (ระบบภายในบริษัท ผู้ใช้ไม่เกินหลักสิบคน) — p95 = 95% ของคำขอเร็วกว่าค่านี้")
    note.font = font(9, italic=True, color="6B7280")
    r += 2
    ws_plan.cell(row=r, column=1, value="รูปแบบการทดสอบ").font = font(11, True, NAVY)
    r += 1
    header(ws_plan, r, ["รูปแบบ", "รูปแบบโหลด"])
    for t, desc in findings.get("profiles", []):
        r += 1
        body(ws_plan.cell(row=r, column=1, value=TYPE_TH.get(t, t)), bold=True)
        body(ws_plan.cell(row=r, column=2, value=desc))
    r += 2
    ws_plan.cell(row=r, column=1, value="สิ่งที่ผู้ใช้จำลองทำ").font = font(11, True, NAVY)
    r += 1
    header(ws_plan, r, ["กลุ่ม", "ลำดับการใช้งาน"])
    for g, desc in findings.get("journeys", []):
        r += 1
        body(ws_plan.cell(row=r, column=1, value=g), bold=True)
        body(ws_plan.cell(row=r, column=2, value=desc))
    r += 2
    ws_plan.cell(row=r, column=1, value="ข้อจำกัด").font = font(11, True, NAVY)
    for line in findings.get("limitations", []):
        r += 1
        c = ws_plan.cell(row=r, column=1, value=f"• {line}")
        c.font = font(10)
        ws_plan.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws_plan.row_dimensions[r].height = 30

    # ---------- ผลรายเส้น API ----------
    title(ws_ep, "ผลรายเส้น API", "เวลาตอบเป็นมิลลิวินาที (ms) คิดจากทุกคำขอ · ผ่าน = p95 ไม่เกินเกณฑ์ และ error ไม่เกินเกณฑ์ (สูตรอ้างแผ่นงาน 'วิธีทดสอบและเกณฑ์')")
    cols = ["รูปแบบ", "เส้น API", "ทำอะไร", "จำนวนคำขอ", "ไม่สำเร็จ", "อัตรา error", "เฉลี่ย", "p50", "p90", "p95", "p99", "สูงสุด", "ผล"]
    header(ws_ep, 4, cols, [10, 30, 32, 11, 10, 10, 9, 9, 9, 9, 9, 10, 10])
    r = 4
    ep_rows = defaultdict(list)  # type -> [row,...]
    for t in TYPES:
        if t not in data:
            continue
        by = defaultdict(list)
        fails = defaultdict(int)
        for ts, name, d, fl, st in data[t]["reqs"]:
            by[name].append(d)
            fails[name] += 1 if fl >= 1 else 0
        for name in sorted(by, key=lambda n: (n.split(":")[0] != "public", n)):
            ds = by[name]
            r += 1
            ep_rows[t].append(r)
            vals = [t, name, ENDPOINT_TH.get(name, ""), len(ds), fails[name], None,
                    sum(ds) / len(ds), pct(ds, 50), pct(ds, 90), pct(ds, 95), pct(ds, 99), max(ds), None]
            for i, v in enumerate(vals, start=1):
                c = ws_ep.cell(row=r, column=i, value=v)
                body(c, "#,##0" if i in (4, 5) else ("#,##0.0" if 7 <= i <= 12 else None), align="right" if i >= 4 else "left")
            ws_ep.cell(row=r, column=6, value=f"=IF(D{r}=0,0,E{r}/D{r})").number_format = "0.00%"
            crit = crit_rows["pdf_p95"] if name.startswith("pdf:") else crit_rows["read_p95"]
            ws_ep.cell(row=r, column=13, value=f'=IF(AND(J{r}<={crit},F{r}<={crit_rows["err"]}),"ผ่าน","ไม่ผ่าน")')
            ws_ep.cell(row=r, column=13).alignment = Alignment(horizontal="center")
    last_ep = r
    ws_ep.conditional_formatting.add(f"M5:M{last_ep}", CellIsRule(operator="equal", formula=['"ผ่าน"'], fill=fill("DCFCE7"), font=Font(name=FONT, bold=True, color="166534")))
    ws_ep.conditional_formatting.add(f"M5:M{last_ep}", CellIsRule(operator="equal", formula=['"ไม่ผ่าน"'], fill=fill("FEE2E2"), font=Font(name=FONT, bold=True, color="B91C1C")))
    ws_ep.freeze_panes = "C5"
    ws_ep.auto_filter.ref = f"A4:M{last_ep}"

    # ---------- ผลตามช่วงเวลา (ทุก 15 วินาที) ----------
    title(ws_st, "ผลตามช่วงเวลา", "แต่ละแถว = 15 วินาที · ผู้ใช้พร้อมกัน = รวมทุกกลุ่ม · ใช้ดูว่าเวลาตอบเปลี่ยนไปอย่างไรเมื่อผู้ใช้เพิ่มขึ้น")
    header(ws_st, 4, ["รูปแบบ", "วินาทีที่", "ผู้ใช้พร้อมกัน", "คำขอ", "คำขอ/วินาที", "p95 (ms)", "error"], [10, 10, 13, 10, 12, 11, 10])
    r = 4
    st_blocks = {}
    for t in TYPES:
        if t not in data or t == "smoke":
            continue
        reqs = data[t]["reqs"]
        if not reqs:
            continue
        t0 = min(x[0] for x in reqs)
        bins = defaultdict(list)
        bfail = defaultdict(int)
        for ts, name, d, fl, st in reqs:
            b = (ts - t0) // 15
            bins[b].append(d)
            bfail[b] += 1 if fl >= 1 else 0
        start = r + 1
        for b in sorted(bins):
            r += 1
            window = [v for ts2, v in data[t]["vus"].items() if t0 + b * 15 <= ts2 < t0 + (b + 1) * 15]
            ds = bins[b]
            vals = [t, int(b * 15), max(window) if window else None, len(ds), None, pct(ds, 95), None]
            for i, v in enumerate(vals, start=1):
                body(ws_st.cell(row=r, column=i, value=v), "#,##0" if i in (2, 3, 4) else ("#,##0.0" if i == 6 else None), align="right" if i > 1 else "left")
            ws_st.cell(row=r, column=5, value=f"=D{r}/15").number_format = "#,##0.0"
            ws_st.cell(row=r, column=7, value=f"=IF(D{r}=0,0,{bfail[b]}/D{r})").number_format = "0.00%"
        st_blocks[t] = (start, r)
        r += 1
    # กราฟ stress: ผู้ใช้ vs p95
    chart_col = 9
    anchor_row = 4
    for t in ("stress", "spike", "load"):
        if t not in st_blocks:
            continue
        a, b = st_blocks[t]
        ch = LineChart()
        ch.title = f"{t}: p95 (ms) เทียบผู้ใช้พร้อมกัน"
        ch.y_axis.title = "p95 (ms)"
        ch.height, ch.width = 7.5, 16
        ch.add_data(Reference(ws_st, min_col=6, min_row=a, max_row=b), titles_from_data=False)
        ch.series[0].tx = SeriesLabel(v="p95 (ms)")
        cats = Reference(ws_st, min_col=2, min_row=a, max_row=b)
        ch.set_categories(cats)
        users = LineChart()
        users.add_data(Reference(ws_st, min_col=3, min_row=a, max_row=b), titles_from_data=False)
        users.y_axis.axId = 200
        users.y_axis.title = "ผู้ใช้"
        users.y_axis.crosses = "max"
        users.series[0].tx = SeriesLabel(v="ผู้ใช้พร้อมกัน")
        show_axes(ch, users)
        ch.legend.position = "b"
        ch += users
        ws_st.add_chart(ch, f"{get_column_letter(chart_col)}{anchor_row}")
        anchor_row += 16

    # ---------- ทรัพยากรเครื่อง ----------
    title(ws_res, "ทรัพยากรเครื่องระหว่างทดสอบ", "เก็บทุก 5 วินาที · CPU ของ API คิดเป็น % ของทั้งเครื่อง (ทุกคอร์รวม = 100%) · ฐานข้อมูล = คอนเทนเนอร์ postgres (100% = 1 คอร์)")
    header(ws_res, 4, ["รูปแบบ", "วินาทีที่", "API CPU %", "API RAM (MB)", "DB CPU %", "DB RAM (MB)"], [10, 10, 12, 13, 11, 13])
    r = 4
    res_blocks = {}
    for t in TYPES:
        if t not in data or not data[t]["res"]:
            continue
        rows = data[t]["res"]
        t0 = rows[0][0]
        start = r + 1
        for at, ac, am, dc, dm in rows:
            r += 1
            vals = [t, round((at - t0).total_seconds()), ac, am, dc, dm]
            for i, v in enumerate(vals, start=1):
                body(ws_res.cell(row=r, column=i, value=v), "#,##0" if i in (2, 4, 6) else ("0.0" if i in (3, 5) else None), align="right" if i > 1 else "left")
        res_blocks[t] = (start, r)
        r += 1
    anchor_row = 4
    for t in ("stress", "spike", "pdf"):
        if t not in res_blocks:
            continue
        a, b = res_blocks[t]
        ch = LineChart()
        ch.title = f"{t}: CPU (%)"
        ch.height, ch.width = 7.5, 16
        ch.add_data(Reference(ws_res, min_col=3, min_row=a, max_row=b), titles_from_data=False)
        ch.add_data(Reference(ws_res, min_col=5, min_row=a, max_row=b), titles_from_data=False)
        ch.series[0].tx = SeriesLabel(v="API")
        ch.series[1].tx = SeriesLabel(v="ฐานข้อมูล")
        ch.set_categories(Reference(ws_res, min_col=2, min_row=a, max_row=b))
        ch.y_axis.title = "CPU %"
        ch.legend.position = "b"
        show_axes(ch)
        ws_res.add_chart(ch, f"H{anchor_row}")
        anchor_row += 16

    # ---------- สรุปผล ----------
    title(ws, f"รายงานผล Performance Test — SystemBOQ{(' · ' + round_label) if round_label else ''}", f"สร้างเมื่อ {thai_date(datetime.now())} · ทดสอบบนเครื่องพัฒนาด้วย k6 · รายละเอียดวิธีทดสอบอยู่ที่แผ่นงาน 'วิธีทดสอบและเกณฑ์'")
    cols = ["รูปแบบ", "ผู้ใช้พร้อมกันสูงสุด", "ระยะเวลา (วินาที)", "จำนวนคำขอ", "ไม่สำเร็จ", "อัตรา error", "คำขอ/วินาที", "p95 ทั้งหมด (ms)", "ช้าสุด (ms)", "เส้นที่ไม่ผ่าน", "ผล"]
    header(ws, 4, cols, [38, 13, 12, 12, 10, 10, 11, 13, 12, 12, 12])
    r = 4
    for t in TYPES:
        if t not in data:
            continue
        s = data[t]["summary"]["metrics"]
        reqs = data[t]["reqs"]
        ds = [x[2] for x in reqs]
        dur = data[t]["summary"]["state"]["testRunDurationMs"] / 1000
        rows = ep_rows[t]
        r += 1
        vals = [TYPE_TH[t], PEAK_PER_GROUP[t] * GROUPS[t], round(dur), None, None, None, None, pct(ds, 95), max(ds) if ds else None, None, None]
        for i, v in enumerate(vals, start=1):
            body(ws.cell(row=r, column=i, value=v), "#,##0" if i in (2, 3, 4, 5, 8, 9, 10) else None, align="right" if i > 1 else "left")
        a, b = rows[0], rows[-1]
        ws.cell(row=r, column=4, value=f"=SUM('ผลรายเส้น API'!D{a}:D{b})").number_format = "#,##0"
        ws.cell(row=r, column=5, value=f"=SUM('ผลรายเส้น API'!E{a}:E{b})").number_format = "#,##0"
        ws.cell(row=r, column=6, value=f"=IF(D{r}=0,0,E{r}/D{r})").number_format = "0.00%"
        ws.cell(row=r, column=7, value=f"=D{r}/C{r}").number_format = "#,##0.0"
        ws.cell(row=r, column=10, value=f"=COUNTIF('ผลรายเส้น API'!M{a}:M{b},\"ไม่ผ่าน\")").number_format = "#,##0"
        ws.cell(row=r, column=11, value=f'=IF(J{r}=0,"ผ่าน","ไม่ผ่าน")')
        ws.cell(row=r, column=11).alignment = Alignment(horizontal="center")
        for col in (8, 9):
            ws.cell(row=r, column=col).number_format = "#,##0.0"
    last = r
    ws.conditional_formatting.add(f"K5:K{last}", CellIsRule(operator="equal", formula=['"ผ่าน"'], fill=fill("DCFCE7"), font=Font(name=FONT, bold=True, color="166534")))
    ws.conditional_formatting.add(f"K5:K{last}", CellIsRule(operator="equal", formula=['"ไม่ผ่าน"'], fill=fill("FEE2E2"), font=Font(name=FONT, bold=True, color="B91C1C")))
    r += 1
    body(ws.cell(row=r, column=1, value="รวม"), bold=True)
    for col, f in ((4, f"=SUM(D5:D{last})"), (5, f"=SUM(E5:E{last})"), (6, f"=IF(D{r}=0,0,E{r}/D{r})")):
        c = ws.cell(row=r, column=col, value=f)
        body(c, "0.00%" if col == 6 else "#,##0", bold=True, align="right")
    r += 2
    ws.cell(row=r, column=1, value="สรุปสั้น").font = font(11, True, NAVY)
    for line in findings.get("summary", []):
        r += 1
        c = ws.cell(row=r, column=1, value=f"• {line}")
        c.font = font(10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=11)
        ws.row_dimensions[r].height = 32
    # กราฟ p95 รายรูปแบบ
    ch = BarChart()
    ch.title = "p95 ของทุกคำขอ (ms) แยกตามรูปแบบ"
    ch.height, ch.width = 8.5, 26
    ch.add_data(Reference(ws, min_col=8, min_row=5, max_row=last), titles_from_data=False)
    ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=last))
    ch.legend = None
    ch.varyColors = False
    ch.y_axis.title = "ms"
    ch.dataLabels = DataLabelList()
    ch.dataLabels.showVal = True
    ch.dataLabels.numFmt = "#,##0"
    ch.dataLabels.showSerName = False
    ch.dataLabels.showCatName = False
    ch.dataLabels.showLegendKey = False
    show_axes(ch)
    ws.add_chart(ch, f"A{r + 2}")

    # ---------- ข้อสังเกต ----------
    title(ws_fd, "ข้อสังเกตและข้อเสนอแนะ", "อ่านจากผลในแผ่นงานอื่น · ความรุนแรง: สูง = ควรแก้ก่อนผู้ใช้เพิ่ม · กลาง = ควรแก้ · ต่ำ = ทราบไว้")
    header(ws_fd, 4, ["#", "หัวข้อ", "สิ่งที่พบ (ตัวเลขจากผลทดสอบ)", "สาเหตุที่น่าจะเป็น", "ข้อเสนอแนะ", "ความรุนแรง"], [5, 26, 50, 40, 46, 12])
    r = 4
    for i, f in enumerate(findings.get("findings", []), start=1):
        r += 1
        for col, v in enumerate([i, f["topic"], f["found"], f["cause"], f["advice"], f["severity"]], start=1):
            body(ws_fd.cell(row=r, column=col, value=v), align="center" if col in (1, 6) else "left")
        ws_fd.row_dimensions[r].height = 75
    if r > 4:
        sev = {"สูง": ("FEE2E2", "B91C1C"), "กลาง": ("FEF3C7", "92400E"), "ต่ำ": ("E0F2FE", "075985")}
        for k, (bg, fg) in sev.items():
            ws_fd.conditional_formatting.add(f"F5:F{r}", CellIsRule(operator="equal", formula=[f'"{k}"'], fill=fill(bg), font=Font(name=FONT, bold=True, color=fg)))

    if "--compare" in opts:
        labels = findings.get("compare_labels", ["รอบก่อน", "รอบนี้"])
        compare_sheet(wb, load_round(Path(opts["--compare"])), data,
                      Path(opts["--first-pdf"]) if "--first-pdf" in opts else None, labels[0], labels[1])

    for s in wb.worksheets:
        s.sheet_view.zoomScale = 100
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"บันทึก {out} · รูปแบบ {list(data)}")


if __name__ == "__main__":
    main()
